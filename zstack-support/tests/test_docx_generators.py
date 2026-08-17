from __future__ import annotations

import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

if sys.version_info < (3, 10):
    raise RuntimeError("DOCX generator tests require Python 3.10 or newer.")


PLUGIN_DIR = Path(__file__).resolve().parents[1]
FAULT_SCRIPT = PLUGIN_DIR / "skills" / "故障报告" / "scripts" / "generate_fault_report.py"
CHANGE_SCRIPT = PLUGIN_DIR / "skills" / "变更方案" / "scripts" / "generate_change_proposal.py"


def load_generator(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load generator: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


fault = load_generator("zstack_fault_generator", FAULT_SCRIPT)
change = load_generator("zstack_change_generator", CHANGE_SCRIPT)

from docx import Document  # noqa: E402


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CORE_TAGS = {
    "{http://purl.org/dc/elements/1.1/}creator",
    "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}lastModifiedBy",
    "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}revision",
    "{http://purl.org/dc/terms/}created",
    "{http://purl.org/dc/terms/}modified",
}


def fault_data() -> dict[str, object]:
    return {
        "audience": "internal",
        "report_title": "云平台资源操作异常故障报告",
        "report_subtitle": "——问题故障报告",
        "project_name": "测试项目",
        "customer_name": "测试客户",
        "fault_impact_scope": "影响一个管理操作，业务范围待确认。",
        "customer_contact": "待补充",
        "reporter": "支持工程师",
        "software_product": "ZStack Cloud",
        "version": "5.3.0",
        "fault_start_time": "2026-07-12 10:00",
        "business_recovery_time": "2026-07-12 10:30",
        "fault_duration": "30 分钟",
        "fault_category": "软件",
        "responsible_department": "技术支持",
        "fault_level": "三级（一般）",
        "fault_description": ["故障现象：资源操作返回异常。", "影响对象：单个资源。"],
        "analysis_process": ["时间：10:05；动作：收集日志；结果：完成。"],
        "root_cause_analysis": ["初步判断：数据状态异常，仍需确认。", "证据边界：当前仅有日志与现场现象。"],
        "improvement_plan": ["改进措施：**增加变更前检查**。", "责任团队：技术支持。"],
    }


def physical_host_reboot_regression_data() -> dict[str, object]:
    data = fault_data()
    data.update(
        {
            "report_title": "物理机异常重启故障报告（格式回归样例）",
            "fault_description": [
                "目标物理机发生异常重启，具体业务影响待确认。",
                "同集群对照节点未出现同类重启现象。",
            ],
            "analysis_process": [
                "排查对象：核对目标节点重启前后的系统与带外记录。",
                "对照对象：比较同集群节点的驱动、固件和运行状态。",
            ],
            "root_cause_analysis": [
                "初步判断：现有证据提示需继续核查 mpt3sas 相关异常，仍需确认。",
                "证据边界：该内容仅用于生成器格式回归，不代表真实客户根因结论。",
            ],
            "improvement_plan": [
                {
                    "title": "取证增强",
                    "prerequisites": ["维护窗口可用"],
                    "actions": ["保持 kdump", "配置 netconsole"],
                    "owner": "待指定",
                    "validation": ["确认 vmcore 及远程日志可落盘"],
                    "rollback": [],
                }
            ],
        }
    )
    return data


def change_data() -> dict[str, object]:
    return {
        "audience": "internal",
        "cover_title": "ZStack运维变更方案",
        "change_title": "测试项目授权导入变更",
        "risk_level": "低",
        "document_date": "2026-07-12",
        "software_info": "ZStack Cloud 5.3.0",
        "hardware_config": "本次变更不涉及硬件调整。",
        "business_info": "影响范围：待确认。",
        "overview": ["变更目标：导入已审批的授权文件。"],
        "change_principles": ["变更前确认文件来源、版本和审批结果。"],
        "overall_flow": ["检查、导入、验证、观察。"],
        "detailed_steps": ["执行动作：按审批文件执行导入。", "验证动作：确认**授权状态**和相关功能。"],
        "risks": ["授权文件与平台版本不匹配可能导致导入失败。"],
        "risk_mitigations": ["执行前由客户或项目负责人确认文件和变更窗口。"],
        "rollback_plan": ["回退动作：恢复客户确认的原授权文件。"],
        "emergency_plan": ["出现非预期影响时立即暂停后续操作并升级处理。"],
        "change_time": "2026-07-12 22:00",
        "maintenance_window": "22:00-23:00",
        "executor": "执行人",
        "executor_phone": "待补充",
        "supervisor": "监督人",
        "supervisor_phone": "待补充",
        "checklist_items": [],
    }


def customer_fault_data(identity_handling: str = "generalized") -> dict[str, object]:
    data = fault_data()
    data["audience"] = "customer"
    data["customer_identity_handling"] = identity_handling
    if identity_handling == "generalized":
        data.update(fault.GENERALIZED_FAULT_IDENTITIES)
    return data


def customer_change_data(identity_handling: str = "generalized") -> dict[str, object]:
    data = change_data()
    data["audience"] = "customer"
    data["customer_identity_handling"] = identity_handling
    return data


def all_document_text(path: Path) -> str:
    document = Document(str(path))
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def rewrite_package(
    source: Path,
    destination: Path,
    replacements: dict[str, bytes] | None = None,
    additions: dict[str, bytes] | None = None,
    removals: set[str] | None = None,
) -> None:
    replacements = replacements or {}
    additions = additions or {}
    removals = removals or set()
    with zipfile.ZipFile(source, "r") as archive_in, zipfile.ZipFile(destination, "w") as archive_out:
        for item in archive_in.infolist():
            if item.filename in removals:
                continue
            archive_out.writestr(item, replacements.get(item.filename, archive_in.read(item.filename)))
        for name, data in additions.items():
            archive_out.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)


def assert_clean_package(test: unittest.TestCase, path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        test.assertIsNone(archive.testzip())
        names = {name.replace("\\", "/").lower() for name in archive.namelist()}
        test.assertNotIn("docprops/custom.xml", names)
        test.assertNotIn("docprops/app.xml", names)
        test.assertFalse(any(name.startswith("docprops/thumbnail.") for name in names))
        test.assertFalse(any(name.startswith("customxml/") for name in names))
        core = ET.fromstring(archive.read("docProps/core.xml"))
        test.assertTrue(CORE_TAGS.isdisjoint({child.tag for child in core}))
        settings = ET.fromstring(archive.read("word/settings.xml"))
        test.assertIsNone(settings.find(f"{{{W_NS}}}docVars"))
        test.assertIsNone(settings.find(f"{{{W_NS}}}rsids"))
        test.assertFalse(any(node.tag.rsplit("}", 1)[-1].casefold() == "docid" for node in settings.iter()))
        settings_bytes = archive.read("word/settings.xml").lower()
        for marker in (b"commondata", b"userid", b"hdid", b"wpscustom"):
            test.assertNotIn(marker, settings_bytes)
        unpacked = b"\n".join(archive.read(name) for name in archive.namelist())
        for secret in (b"userId", b"hdid", b"commondata", b"KSOTemplateDocerSaveRecord"):
            test.assertNotIn(secret, unpacked)
        for name in (item for item in archive.namelist() if item.endswith(".xml")):
            root = ET.fromstring(archive.read(name))
            for element in root.iter():
                test.assertFalse(
                    any(
                        attribute.rsplit("}", 1)[-1].casefold().startswith("rsid")
                        or attribute.rsplit("}", 1)[-1].casefold() in {"durableid", "paraid", "textid"}
                        for attribute in element.attrib
                    ),
                    f"tracking attribute remains in {name}",
                )


class ValidationTests(unittest.TestCase):
    def test_missing_private_dependency_has_actionable_error(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            environment = os.environ.copy()
            environment.pop("PYTHONPATH", None)
            environment["ZSTACK_SUPPORT_PYTHONPATH"] = directory
            result = subprocess.run(
                [sys.executable, "-S", str(FAULT_SCRIPT), "--print-template"],
                env=environment,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                check=False,
            )
            self.assertEqual(result.returncode, 2)
            self.assertIn("pip install --target", result.stderr)
            self.assertIn("requirements.txt", result.stderr)

    def test_empty_missing_wrong_type_and_unknown_fields_fail(self) -> None:
        for validator, factory in ((fault.validate_input, fault_data), (change.validate_input, change_data)):
            with self.subTest(validator=validator.__module__, case="empty"):
                with self.assertRaisesRegex(ValueError, "Missing required"):
                    validator({})
            missing = factory()
            missing.pop(next(iter(missing)))
            with self.subTest(validator=validator.__module__, case="missing"):
                with self.assertRaises(ValueError):
                    validator(missing)
            wrong = factory()
            wrong[next(iter(wrong))] = 7
            with self.subTest(validator=validator.__module__, case="type"):
                with self.assertRaisesRegex(ValueError, "must be a string"):
                    validator(wrong)
            unknown = factory()
            unknown["unexpected"] = "value"
            with self.subTest(validator=validator.__module__, case="unknown"):
                with self.assertRaisesRegex(ValueError, "Unknown input field"):
                    validator(unknown)

    def test_empty_value_and_unbounded_root_cause_fail(self) -> None:
        data = fault_data()
        data["customer_name"] = "  "
        with self.assertRaisesRegex(ValueError, "must not be empty"):
            fault.validate_input(data)
        data = fault_data()
        data["root_cause_analysis"] = ["数据异常导致故障。"]
        with self.assertRaisesRegex(ValueError, "evidence boundary"):
            fault.validate_input(data)

    def test_explicit_pending_values_are_allowed(self) -> None:
        data = fault_data()
        data["customer_contact"] = "待补充"
        self.assertEqual(fault.validate_input(data)["customer_contact"], "待补充")

    def test_checklist_booleans_are_strict(self) -> None:
        data = change_data()
        data["mark_unmatched_checklist_as_no"] = "false"
        with self.assertRaisesRegex(ValueError, "JSON boolean"):
            change.validate_input(data)

    def test_audience_and_customer_identity_contract(self) -> None:
        for validator, factory in ((fault.validate_input, fault_data), (change.validate_input, change_data)):
            invalid = factory()
            invalid["audience"] = "external"
            with self.subTest(validator=validator.__module__, case="audience"):
                with self.assertRaisesRegex(ValueError, "internal.*customer"):
                    validator(invalid)
            customer = factory()
            customer["audience"] = "customer"
            with self.subTest(validator=validator.__module__, case="identity-required"):
                with self.assertRaisesRegex(ValueError, "customer_identity_handling"):
                    validator(customer)

        self.assertEqual(fault.validate_input(customer_fault_data())["audience"], "customer")
        self.assertEqual(change.validate_input(customer_change_data())["audience"], "customer")
        confirmed = customer_fault_data("confirmed_same_customer")
        self.assertEqual(
            fault.validate_input(confirmed)["customer_identity_handling"], "confirmed_same_customer"
        )
        generalized = customer_fault_data()
        generalized["customer_name"] = "某客户"
        with self.assertRaisesRegex(ValueError, "exactly"):
            fault.validate_input(generalized)

    def test_customer_output_rejects_internal_links_endpoints_and_raw_ids(self) -> None:
        cases = {
            "bbs": "https://bbs.internal.example/topic/1",
            "jira": "https://support.atlassian.net/jira/browse/CASE-1",
            "confluence": "https://confluence.example.com/display/KB",
            "private-ip": "受控端点 192.168.10.12:3340",
            "public-ip": "公开 IP 8.8.8.8",
            "ipv6": "IPv6 2001:4860:4860::8888",
            "uuid": "资源 " + "123e4567" + "-e89b-12d3-a456-426614174000",
            "compact-guid": "资源 " + "123e4567e89b12d3a456426614174000",
            "integer-ipv4": "端点 2130706433",
            "hex-ipv4": "端点 0x7f000001",
            "octal-ipv4": "端点 017700000001",
            "mixed-ipv4": "端点 0x7f.00.0.01",
            "leading-zero-ipv4": "端点 00127.000.000.001",
            "abbreviated-ipv4-url": "http://127.1",
            "abbreviated-ipv4": "127.0.1",
            "abbreviated-octal-ipv4": "0177.1",
            "encoded-link": "https%3A%2F%2Fjira.internal.example%2Fbrowse%2FTIC-1",
            "encoded-host": "https://%6a%69%72%61.internal.example/browse/TIC-1",
            "ticket": "内部工单 TIC-5786",
            "bug": "内部缺陷 BUG-123",
            "suggestion": "内部建议 SUG-123",
            "relative-browse": "/browse/TIC-123",
            "relative-forum": "forum.php?mod=viewthread&tid=42",
            "relative-confluence": "confluence/pages/viewpage.action?pageId=1",
            "short-basic": "Authorization: " + "Basic " + "dTpw",
            "short-bearer": "Authorization: " + "Bearer " + "x",
            "encrypted-private-key": "-----BEGIN " + "ENCRYPTED PRIVATE KEY-----",
            "pgp-private-key": "-----BEGIN PGP " + "PRIVATE KEY BLOCK-----",
        }
        for name, forbidden in cases.items():
            for validator, data, field in (
                (change.validate_input, customer_change_data(), "overview"),
                (fault.validate_input, customer_fault_data(), "fault_description"),
            ):
                data[field] = [forbidden]
                with self.subTest(case=name, validator=validator.__module__):
                    with self.assertRaises(ValueError):
                        validator(data)

    def test_customer_output_allows_public_non_support_links(self) -> None:
        safe_values = (
            "Public reference: https://github.com/zstackio/zstack/issues/123",
            "version 1.2.3.4",
            "build 20260712",
            "checksum d41d8cd98f00b204e9800998ecf8427e",
            "Basic authentication is disabled",
            "token=disabled",
            "literal 0x1",
        )
        for value in safe_values:
            change_input = customer_change_data()
            change_input["overview"] = [value]
            fault_input = customer_fault_data()
            fault_input["fault_description"] = [value]
            with self.subTest(value=value):
                self.assertEqual(change.validate_input(change_input)["audience"], "customer")
                self.assertEqual(fault.validate_input(fault_input)["audience"], "customer")
        data = change_data()
        data["checklist_decisions"] = [{"operation": "删除区域", "involved": "false"}]
        with self.assertRaisesRegex(ValueError, "JSON boolean"):
            change.validate_input(data)

    def test_utf8_bom_load(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fault.json"
            path.write_text(json.dumps(fault_data(), ensure_ascii=False), encoding="utf-8-sig")
            self.assertEqual(fault.load_json(path)["project_name"], "测试项目")

    def test_duplicate_json_keys_are_rejected_at_any_depth(self) -> None:
        samples = (
            '{"audience":"internal","audience":"customer"}',
            '{"outer":{"key":1,"key":2}}',
        )
        with tempfile.TemporaryDirectory() as directory:
            for index, sample in enumerate(samples):
                path = Path(directory) / f"duplicate-{index}.json"
                path.write_text(sample, encoding="utf-8-sig")
                for loader in (fault.load_json, change.load_json):
                    with self.subTest(index=index, loader=loader.__module__):
                        with self.assertRaisesRegex(ValueError, "duplicate key"):
                            loader(path)

    def test_internal_and_customer_outputs_reject_obvious_secrets(self) -> None:
        secrets = (
            "Authorization: " + "Bearer " + "abcdefghijklmnop",
            "password=NotARealPassword123",
            "-----BEGIN " + "PRIVATE KEY-----",
            "https://user:password@example.com/path",
            "ghp_" + "A" * 24,
        )
        for audience in ("internal", "customer"):
            for secret in secrets:
                change_input = change_data() if audience == "internal" else customer_change_data()
                fault_input = fault_data() if audience == "internal" else customer_fault_data()
                change_input["overview"] = [secret]
                fault_input["fault_description"] = [secret]
                for validator, data in ((change.validate_input, change_input), (fault.validate_input, fault_input)):
                    with self.subTest(audience=audience, secret=secret[:20], validator=validator.__module__):
                        with self.assertRaises(ValueError):
                            validator(data)


class ChecklistTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.rows = change.checklist_rows(Document(str(change.DEFAULT_TEMPLATE)))

    def test_stable_row_ids_and_level_contract(self) -> None:
        ids = [row["row_id"] for row in self.rows]
        self.assertEqual(len(ids), 65)
        self.assertEqual(len(ids), len(set(ids)))
        self.assertTrue(all(re.fullmatch(r"cl-[0-9a-f]{12}", row_id) for row_id in ids))
        self.assertEqual(
            {row["module_id"] for row in self.rows},
            {module_id for module_id, _, _ in change.CHECKLIST_MODULE_CONTRACT},
        )
        for row in self.rows:
            self.assertEqual(
                row["row_id"],
                change.stable_row_id(row["module_id"], row["operation"], row["impact"], row["level"]),
            )
        osd = next(row for row in self.rows if row["operation"] == "OSD数据不均衡调整")
        self.assertEqual(osd["level"], "中")

    def test_module_and_level_contract_do_not_use_name_substrings(self) -> None:
        document = Document(str(change.DEFAULT_TEMPLATE))
        table = change.unique_table_by_first_cell(document, change.CHECKLIST_TITLE)
        table.rows[2].cells[0].text = "自定义高风险模块"
        with self.assertRaisesRegex(ValueError, "unknown module contract"):
            change.checklist_rows(document)

        document = Document(str(change.DEFAULT_TEMPLATE))
        table = change.unique_table_by_first_cell(document, change.CHECKLIST_TITLE)
        table.rows[2].cells[3].text = "中"
        with self.assertRaisesRegex(ValueError, "module contract"):
            change.checklist_rows(document)

    def test_duplicate_operation_requires_disambiguation(self) -> None:
        with self.assertRaisesRegex(ValueError, "ambiguous"):
            change.resolve_selector("交换机变更调整", self.rows, "test")
        duplicate = [row for row in self.rows if row["operation"] == "交换机变更调整"]
        selected = change.resolve_selector({"row_id": duplicate[0]["row_id"]}, self.rows, "test")
        self.assertEqual(selected["impact"], duplicate[0]["impact"])

    def test_unmatched_checklist_item_fails_without_output(self) -> None:
        data = change_data()
        data["checklist_items"] = ["不存在的操作"]
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output.docx"
            with self.assertRaisesRegex(ValueError, "did not match"):
                change.generate(data, change.DEFAULT_TEMPLATE, output)
            self.assertFalse(output.exists())


class PathAndAtomicityTests(unittest.TestCase):
    def test_input_output_and_template_collisions_fail(self) -> None:
        input_path = Path("input.json").resolve()
        template = change.DEFAULT_TEMPLATE.resolve()
        with self.assertRaisesRegex(ValueError, "input JSON"):
            change.validate_paths(input_path, template, input_path)
        with self.assertRaisesRegex(ValueError, "template"):
            change.validate_paths(input_path, template, template)

    def test_template_drift_preserves_existing_output(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            drifted = directory_path / "drifted.docx"
            shutil.copyfile(change.DEFAULT_TEMPLATE, drifted)
            document = Document(str(drifted))
            paragraph = next(item for item in document.paragraphs if item.text.strip() == "XX变更")
            paragraph.text = "DRIFTED"
            document.save(str(drifted))
            output = directory_path / "output.docx"
            output.write_bytes(b"existing-output")
            with self.assertRaisesRegex(ValueError, "XX变更"):
                change.generate(change_data(), drifted, output)
            self.assertEqual(output.read_bytes(), b"existing-output")

    def test_cli_rejects_path_collision_before_writing(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory) / "input.json"
            original = json.dumps(change_data(), ensure_ascii=False).encode("utf-8")
            input_path.write_bytes(original)
            result = subprocess.run(
                [sys.executable, str(CHANGE_SCRIPT), str(input_path), "--out", str(input_path)],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                check=False,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("input JSON", result.stderr)
            self.assertEqual(input_path.read_bytes(), original)

    def test_customer_post_render_scan_is_atomic(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            template = directory_path / "template-with-internal-link.docx"
            document = Document(str(change.DEFAULT_TEMPLATE))
            overview_heading = next(item for item in document.paragraphs if item.text.strip() == "变更概述")
            overview_heading.insert_paragraph_before("https://jira.internal.example/browse/CASE-1")
            document.save(str(template))
            output = directory_path / "customer.docx"
            output.write_bytes(b"existing-output")
            with self.assertRaisesRegex(ValueError, "internal"):
                change.generate(customer_change_data(), template, output)
            self.assertEqual(output.read_bytes(), b"existing-output")

    def test_fault_customer_post_render_scan_is_atomic(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            template = directory_path / "fault-template-with-internal-link.docx"
            document = Document(str(fault.DEFAULT_TEMPLATE))
            cover = next(item for item in document.paragraphs if item.text.strip() == "云平台数据库残留XXXX")
            cover.insert_paragraph_before("/browse/TIC-123")
            document.save(str(template))
            output = directory_path / "customer-fault.docx"
            output.write_bytes(b"existing-output")
            with self.assertRaisesRegex(ValueError, "internal"):
                fault.generate(customer_fault_data(), template, output)
            self.assertEqual(output.read_bytes(), b"existing-output")


class PackageSecurityTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_default_templates_pass_relationship_graph_scan(self) -> None:
        for module in (fault, change):
            with zipfile.ZipFile(module.DEFAULT_TEMPLATE) as archive:
                module.validate_customer_package(archive)

    def test_customer_rejects_media_not_in_bundled_hash_allowlist(self) -> None:
        with zipfile.ZipFile(change.DEFAULT_TEMPLATE) as other:
            replacement = other.read("word/media/image1.png")
        mutated = self.directory / "tampered-media.docx"
        rewrite_package(
            fault.DEFAULT_TEMPLATE,
            mutated,
            replacements={"word/media/image1.png": replacement},
        )
        with zipfile.ZipFile(mutated) as archive:
            with self.assertRaisesRegex(ValueError, "media part"):
                fault.validate_customer_package(archive)

    def test_customer_rejects_unreferenced_hidden_style(self) -> None:
        with zipfile.ZipFile(fault.DEFAULT_TEMPLATE) as archive:
            styles = ET.fromstring(archive.read("word/styles.xml"))
        style = ET.SubElement(styles, f"{{{W_NS}}}style", {f"{{{W_NS}}}type": "character", f"{{{W_NS}}}styleId": "HiddenPayload"})
        run_properties = ET.SubElement(style, f"{{{W_NS}}}rPr")
        ET.SubElement(run_properties, f"{{{W_NS}}}vanish")
        mutated = self.directory / "hidden-style.docx"
        rewrite_package(
            fault.DEFAULT_TEMPLATE,
            mutated,
            replacements={"word/styles.xml": ET.tostring(styles, encoding="utf-8", xml_declaration=True)},
        )
        with zipfile.ZipFile(mutated) as archive:
            with self.assertRaisesRegex(ValueError, "hidden-text style"):
                fault.validate_customer_package(archive)

    def test_customer_rejects_hidden_doc_defaults(self) -> None:
        with zipfile.ZipFile(fault.DEFAULT_TEMPLATE) as archive:
            styles = ET.fromstring(archive.read("word/styles.xml"))
        defaults = ET.SubElement(styles, f"{{{W_NS}}}docDefaults")
        run_defaults = ET.SubElement(defaults, f"{{{W_NS}}}rPrDefault")
        run_properties = ET.SubElement(run_defaults, f"{{{W_NS}}}rPr")
        ET.SubElement(run_properties, f"{{{W_NS}}}vanish")
        mutated = self.directory / "hidden-doc-defaults.docx"
        rewrite_package(
            fault.DEFAULT_TEMPLATE,
            mutated,
            replacements={"word/styles.xml": ET.tostring(styles, encoding="utf-8", xml_declaration=True)},
        )
        with zipfile.ZipFile(mutated) as archive:
            with self.assertRaisesRegex(ValueError, "hidden-text"):
                fault.validate_customer_package(archive)

    def test_customer_rejects_unknown_relationship_payload(self) -> None:
        with zipfile.ZipFile(fault.DEFAULT_TEMPLATE) as archive:
            relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
            content_types = ET.fromstring(archive.read("[Content_Types].xml"))
        ET.SubElement(
            relationships,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            {
                "Id": "rIdUnknownPayload",
                "Type": "https://example.invalid/relationships/opaque-payload",
                "Target": "payload.bin",
            },
        )
        ET.SubElement(
            content_types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            {"PartName": "/word/payload.bin", "ContentType": "application/octet-stream"},
        )
        mutated = self.directory / "unknown-relationship-payload.docx"
        rewrite_package(
            fault.DEFAULT_TEMPLATE,
            mutated,
            replacements={
                "word/_rels/document.xml.rels": ET.tostring(
                    relationships, encoding="utf-8", xml_declaration=True
                ),
                "[Content_Types].xml": ET.tostring(content_types, encoding="utf-8", xml_declaration=True),
            },
            additions={"word/payload.bin": b"opaque payload"},
        )
        with zipfile.ZipFile(mutated) as archive:
            with self.assertRaisesRegex(ValueError, "unapproved relationship type"):
                fault.validate_customer_package(archive)

    def test_customer_rejects_active_external_hyperlink_schemes(self) -> None:
        for index, target in enumerate(("file:C:/secret.txt", "javascript:alert(1)", "ms-word:ofe|u|file")):
            with zipfile.ZipFile(fault.DEFAULT_TEMPLATE) as archive:
                relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
            ET.SubElement(
                relationships,
                "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                {
                    "Id": f"rIdActiveLink{index}",
                    "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                    "Target": target,
                    "TargetMode": "External",
                },
            )
            mutated = self.directory / f"active-link-{index}.docx"
            rewrite_package(
                fault.DEFAULT_TEMPLATE,
                mutated,
                replacements={
                    "word/_rels/document.xml.rels": ET.tostring(
                        relationships, encoding="utf-8", xml_declaration=True
                    )
                },
            )
            with zipfile.ZipFile(mutated) as archive:
                with self.subTest(target=target):
                    with self.assertRaisesRegex(ValueError, "unapproved URI scheme"):
                        fault.validate_customer_package(archive)

    def test_customer_rejects_external_image_and_attached_template_relationships(self) -> None:
        for role, target in (
            ("image", "https://example.com/image.png"),
            ("attachedTemplate", "https://example.com/template.dotx"),
            ("hyperlink", "https%3A%2F%2Fjira.internal.example%2Fbrowse%2FTIC-1"),
        ):
            with zipfile.ZipFile(fault.DEFAULT_TEMPLATE) as archive:
                relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
            ET.SubElement(
                relationships,
                "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                {
                    "Id": f"rIdMalicious{role}",
                    "Type": f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{role}",
                    "Target": target,
                    "TargetMode": "External",
                },
            )
            mutated = self.directory / f"external-{role}.docx"
            rewrite_package(
                fault.DEFAULT_TEMPLATE,
                mutated,
                replacements={
                    "word/_rels/document.xml.rels": ET.tostring(
                        relationships, encoding="utf-8", xml_declaration=True
                    )
                },
            )
            with zipfile.ZipFile(mutated) as archive:
                with self.subTest(role=role):
                    with self.assertRaises(ValueError):
                        fault.validate_customer_package(archive)

    def test_orphan_parts_are_scanned_for_secrets_and_rejected_for_customer(self) -> None:
        secret_orphan = self.directory / "secret-orphan.docx"
        rewrite_package(
            fault.DEFAULT_TEMPLATE,
            secret_orphan,
            additions={
                "word/nonstandard-payload.xml": (
                    b"<payload>Authorization: " + b"Bearer " + b"abcdefghijklmnop</payload>"
                )
            },
        )
        with self.assertRaisesRegex(ValueError, "authorization credential"):
            fault.validate_package(secret_orphan, "internal")

        benign_orphan = self.directory / "benign-orphan.docx"
        rewrite_package(
            fault.DEFAULT_TEMPLATE,
            benign_orphan,
            additions={"word/nonstandard-payload.xml": b"<payload>benign</payload>"},
        )
        with zipfile.ZipFile(benign_orphan) as archive:
            with self.assertRaisesRegex(ValueError, "unreferenced or unknown"):
                fault.validate_customer_package(archive)

    def test_customer_scans_reachable_style_and_numbering_attributes(self) -> None:
        with zipfile.ZipFile(fault.DEFAULT_TEMPLATE) as archive:
            numbering = ET.fromstring(archive.read("word/numbering.xml"))
        first = next(iter(numbering.iter()))
        first.set("payload", "0x7f000001")
        mutated = self.directory / "attribute-payload.docx"
        rewrite_package(
            fault.DEFAULT_TEMPLATE,
            mutated,
            replacements={"word/numbering.xml": ET.tostring(numbering, encoding="utf-8", xml_declaration=True)},
        )
        with zipfile.ZipFile(mutated) as archive:
            with self.assertRaisesRegex(ValueError, "integer-form IPv4"):
                fault.validate_customer_package(archive)

    def test_customer_rejects_encoded_field_code_link(self) -> None:
        with zipfile.ZipFile(fault.DEFAULT_TEMPLATE) as archive:
            document = ET.fromstring(archive.read("word/document.xml"))
        body = next(node for node in document if node.tag.rsplit("}", 1)[-1] == "body")
        paragraph = ET.Element(f"{{{W_NS}}}p")
        run = ET.SubElement(paragraph, f"{{{W_NS}}}r")
        instruction = ET.SubElement(run, f"{{{W_NS}}}instrText")
        instruction.text = ' HYPERLINK "https%3A%2F%2Fconfluence.internal.example%2Fpages%2F1" '
        body.insert(0, paragraph)
        mutated = self.directory / "field-code-link.docx"
        rewrite_package(
            fault.DEFAULT_TEMPLATE,
            mutated,
            replacements={"word/document.xml": ET.tostring(document, encoding="utf-8", xml_declaration=True)},
        )
        with zipfile.ZipFile(mutated) as archive:
            with self.assertRaises(ValueError):
                fault.validate_customer_package(archive)

    def test_customer_scans_header_by_relationship_after_part_is_renamed(self) -> None:
        for module in (fault, change):
            original_header_name = None
            with zipfile.ZipFile(module.DEFAULT_TEMPLATE) as archive:
                relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
                content_types = ET.fromstring(archive.read("[Content_Types].xml"))
                header_relationship = next(
                    (
                        relationship
                        for relationship in relationships
                        if relationship.get("Type", "").rsplit("/", 1)[-1].casefold() == "header"
                    ),
                    None,
                )
                if header_relationship is None:
                    header_relationship = ET.SubElement(
                        relationships,
                        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                        {
                            "Id": "rIdRenamedHeader",
                            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header",
                        },
                    )
                    header = ET.Element(f"{{{W_NS}}}hdr")
                else:
                    original_header_name = f"word/{header_relationship.get('Target')}"
                    header = ET.fromstring(archive.read(original_header_name))
                header_relationship.set("Target", "hdrSecret.xml")
                paragraph = ET.SubElement(header, f"{{{W_NS}}}p")
                run = ET.SubElement(paragraph, f"{{{W_NS}}}r")
                ET.SubElement(run, f"{{{W_NS}}}t").text = "10.20.30.40"

                original_part_name = f"/{original_header_name}" if original_header_name else None
                header_override = next(
                    (
                        item
                        for item in content_types
                        if item.get("PartName") == original_part_name
                    ),
                    None,
                )
                if header_override is None:
                    header_override = ET.SubElement(
                        content_types,
                        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                        {
                            "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml",
                        },
                    )
                header_override.set("PartName", "/word/hdrSecret.xml")

            mutated = self.directory / f"{module.__name__}-renamed-header.docx"
            rewrite_package(
                module.DEFAULT_TEMPLATE,
                mutated,
                replacements={
                    "word/_rels/document.xml.rels": ET.tostring(
                        relationships, encoding="utf-8", xml_declaration=True
                    ),
                    "[Content_Types].xml": ET.tostring(
                        content_types, encoding="utf-8", xml_declaration=True
                    ),
                },
                additions={
                    "word/hdrSecret.xml": ET.tostring(header, encoding="utf-8", xml_declaration=True)
                },
                removals={original_header_name} if original_header_name else set(),
            )
            with zipfile.ZipFile(mutated) as archive:
                with self.subTest(module=module.__name__):
                    with self.assertRaisesRegex(ValueError, "(?:raw IP|IPv4)"):
                        module.validate_customer_package(archive)

    def test_thumbnail_is_removed_for_internal_output(self) -> None:
        with zipfile.ZipFile(fault.DEFAULT_TEMPLATE) as archive:
            root_relationships = ET.fromstring(archive.read("_rels/.rels"))
            thumbnail = archive.read("word/media/image1.png")
        ET.SubElement(
            root_relationships,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            {
                "Id": "rIdThumbnailPayload",
                "Type": "http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail",
                "Target": "docProps/thumbnail.png",
            },
        )
        template = self.directory / "thumbnail-template.docx"
        rewrite_package(
            fault.DEFAULT_TEMPLATE,
            template,
            replacements={"_rels/.rels": ET.tostring(root_relationships, encoding="utf-8", xml_declaration=True)},
            additions={"docProps/thumbnail.png": thumbnail},
        )
        output = self.directory / "internal-without-thumbnail.docx"
        fault.generate(fault_data(), template, output)
        with zipfile.ZipFile(output) as archive:
            self.assertFalse(any(name.casefold().startswith("docprops/thumbnail.") for name in archive.namelist()))
            relationships = archive.read("_rels/.rels").lower()
            self.assertNotIn(b"thumbnail", relationships)


class GenerationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_complete_fault_generation(self) -> None:
        output = self.directory / "fault.docx"
        fault.generate(fault_data(), fault.DEFAULT_TEMPLATE, output)
        document = Document(str(output))
        self.assertEqual(next(paragraph.text for paragraph in document.paragraphs if "资源操作异常" in paragraph.text), fault_data()["report_title"])
        text = all_document_text(output)
        self.assertNotIn("云平台数据库残留XXXX", text)
        self.assertNotIn("故障情况描述：\n\n\n具体表现为：", text)
        audit = fault.audit_format_structure(output, fault.validate_input(fault_data()))
        self.assertEqual(audit["ordinary_all_bold_paragraphs"], 0)
        self.assertGreater(audit["mixed_bold_paragraphs"], 0)
        assert_clean_package(self, output)

    def test_physical_host_reboot_format_regression(self) -> None:
        data = physical_host_reboot_regression_data()
        output = self.directory / "physical-host-reboot-regression.docx"
        fault.generate(data, fault.DEFAULT_TEMPLATE, output)
        text = all_document_text(output)
        self.assertIn("同集群对照节点未出现同类重启现象", text)
        self.assertIn("mpt3sas", text)
        self.assertNotIn("OFED", text)
        audit = fault.audit_format_structure(output, fault.validate_input(data))
        self.assertEqual(audit["ordinary_all_bold_paragraphs"], 0)
        self.assertEqual(audit["duplicate_number_paragraphs"], 0)
        self.assertGreaterEqual(audit["native_numbered_paragraphs"], 2)
        assert_clean_package(self, output)

    def test_complete_generalized_customer_fault_generation(self) -> None:
        output = self.directory / "customer-fault.docx"
        fault.generate(customer_fault_data(), fault.DEFAULT_TEMPLATE, output)
        text = all_document_text(output)
        for value in fault.GENERALIZED_FAULT_IDENTITIES.values():
            self.assertIn(value, text)
        assert_clean_package(self, output)

    def test_complete_generalized_customer_change_generation(self) -> None:
        output = self.directory / "customer-change.docx"
        change.generate(customer_change_data(), change.DEFAULT_TEMPLATE, output)
        self.assertTrue(output.is_file())
        change.validate_package(output, "customer")

    def test_complete_change_generation_and_checklist_mark(self) -> None:
        rows = change.checklist_rows(Document(str(change.DEFAULT_TEMPLATE)))
        selected = rows[0]
        data = change_data()
        data["checklist_items"] = [{"row_id": selected["row_id"]}]
        output = self.directory / "change.docx"
        change.generate(data, change.DEFAULT_TEMPLATE, output)
        document = Document(str(output))
        checklist = change.unique_table_by_first_cell(document, change.CHECKLIST_TITLE)
        self.assertEqual(checklist.rows[selected["row_index"]].cells[4].text, "是")
        self.assertEqual(checklist.rows[rows[1]["row_index"]].cells[4].text, "")
        text = all_document_text(output)
        for sample in ("2021-11-10", "XXX客户", "XXXXX", "XX变更"):
            self.assertNotIn(sample, text)
        assert_clean_package(self, output)

    def test_template_and_outputs_are_metadata_clean(self) -> None:
        assert_clean_package(self, fault.DEFAULT_TEMPLATE)
        assert_clean_package(self, change.DEFAULT_TEMPLATE)
        fault_output = self.directory / "fault.docx"
        change_output = self.directory / "change.docx"
        fault.generate(fault_data(), fault.DEFAULT_TEMPLATE, fault_output)
        change.generate(change_data(), change.DEFAULT_TEMPLATE, change_output)
        assert_clean_package(self, fault_output)
        assert_clean_package(self, change_output)

    def test_toc_dirty_and_update_fields(self) -> None:
        output = self.directory / "change.docx"
        change.generate(change_data(), change.DEFAULT_TEMPLATE, output)
        with zipfile.ZipFile(output) as archive:
            settings = ET.fromstring(archive.read("word/settings.xml"))
            update_fields = settings.find(f"{{{W_NS}}}updateFields")
            self.assertIsNotNone(update_fields)
            self.assertEqual(update_fields.get(f"{{{W_NS}}}val"), "true")
            document = ET.fromstring(archive.read("word/document.xml"))
            begin_fields = [
                node
                for node in document.iter(f"{{{W_NS}}}fldChar")
                if node.get(f"{{{W_NS}}}fldCharType") == "begin"
            ]
            self.assertGreater(len(begin_fields), 0)
            self.assertTrue(all(node.get(f"{{{W_NS}}}dirty") == "true" for node in begin_fields))

    def test_run_format_and_table_alignment_are_preserved(self) -> None:
        template = Document(str(change.DEFAULT_TEMPLATE))
        template_cover = next(paragraph for paragraph in template.paragraphs if paragraph.text.strip() == "ZStack运维变更方案")
        template_size = template_cover.runs[0].font.size
        template_bold = template_cover.runs[0].bold
        template_normal_size = template.styles["Normal"].font.size
        config = change.unique_table_by_first_cell(template, "软件信息")
        template_alignment = config.rows[0].cells[1].paragraphs[0].alignment
        template_vertical = config.rows[0].cells[1].vertical_alignment

        output = self.directory / "change.docx"
        change.generate(change_data(), change.DEFAULT_TEMPLATE, output)
        rendered = Document(str(output))
        rendered_cover = next(paragraph for paragraph in rendered.paragraphs if paragraph.text.strip() == "ZStack运维变更方案")
        self.assertEqual(rendered_cover.runs[0].font.size, template_size)
        self.assertEqual(rendered_cover.runs[0].bold, template_bold)
        self.assertEqual(rendered.styles["Normal"].font.size, template_normal_size)
        rendered_config = change.unique_table_by_first_cell(rendered, "软件信息")
        self.assertEqual(rendered_config.rows[0].cells[1].paragraphs[0].alignment, template_alignment)
        self.assertEqual(rendered_config.rows[0].cells[1].vertical_alignment, template_vertical)
        action = next(paragraph for paragraph in rendered.paragraphs if paragraph.text.startswith("执行动作："))
        self.assertTrue(action.runs)
        self.assertTrue(all(run.bold is True for run in action.runs if run.text))

        fault_output = self.directory / "fault-format.docx"
        fault.generate(fault_data(), fault.DEFAULT_TEMPLATE, fault_output)
        fault_rendered = Document(str(fault_output))
        info_table = fault.unique_table_by_first_cell(fault_rendered, "故障基本信息")
        for label in ("故障发生时间", "业务恢复时间", "故障级别"):
            _, _, target = fault.label_target(info_table, label)
            self.assertTrue(all(run.bold is True for run in target.paragraphs[0].runs if run.text))
        root_table = fault.unique_table_by_first_cell(fault_rendered, "原因分析")
        conclusion = root_table.rows[1].cells[0].paragraphs[0]
        conclusion_runs = [run for run in conclusion.runs if run.text]
        self.assertTrue(conclusion_runs[0].bold)
        self.assertTrue(conclusion_runs[0].text.endswith("："))
        self.assertTrue(all(run.bold is False for run in conclusion_runs[1:]))

    def test_fault_label_span_and_line_bold_are_independent(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        fault.write_rich_text(paragraph, "执行动作：先执行**关键步骤**再验证")
        self.assertEqual(paragraph.text, "执行动作：先执行关键步骤再验证")
        runs = [(run.text, run.bold) for run in paragraph.runs if run.text]
        self.assertEqual(
            runs,
            [("执行动作：", True), ("先执行", False), ("关键步骤", True), ("再验证", False)],
        )

        fault.write_rich_text(paragraph, "2、原因：普通正文")
        self.assertEqual(
            [(run.text, run.bold) for run in paragraph.runs if run.text],
            [("2、", False), ("原因：", True), ("普通正文", False)],
        )

        fault.write_rich_text(paragraph, "整行需要加粗", line_bold=True)
        self.assertTrue(all(run.bold is True for run in paragraph.runs if run.text))

    def test_fault_plain_text_does_not_inherit_bold_prototype(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        prototype = paragraph.add_run("模板粗体")
        prototype.bold = True
        fault.write_rich_text(paragraph, "普通正文")
        self.assertEqual(paragraph.text, "普通正文")
        self.assertTrue(all(run.bold is False for run in paragraph.runs if run.text))

    def test_change_legacy_bold_behavior_remains_unchanged(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        change.write_rich_text(paragraph, "执行动作：先执行**关键步骤**再验证")
        self.assertTrue(all(run.bold is True for run in paragraph.runs if run.text))

    def test_fault_numbering_is_native_and_idempotent(self) -> None:
        data = fault_data()
        data["fault_description"] = [
            "普通条目 A",
            "普通条目 B",
            "1. 手工编号",
            "2、中文编号",
            "**3.** Markdown 编号",
            "4. 4. 重复编号输入",
        ]
        output = self.directory / "fault-numbering.docx"
        fault.generate(data, fault.DEFAULT_TEMPLATE, output)
        document = Document(str(output))
        info_table = fault.unique_table_by_first_cell(document, "故障基本信息")
        _, _, target = fault.description_target(info_table)
        paragraphs = [paragraph for paragraph in target.paragraphs if paragraph.text.strip()]
        self.assertEqual(
            [paragraph.text for paragraph in paragraphs],
            ["普通条目 A", "普通条目 B", "1. 手工编号", "2、中文编号", "3. Markdown 编号", "4. 重复编号输入"],
        )
        self.assertIsNotNone(fault.paragraph_number_id(paragraphs[0]))
        self.assertEqual(fault.paragraph_number_id(paragraphs[0]), fault.paragraph_number_id(paragraphs[1]))
        self.assertTrue(all(fault.paragraph_number_id(paragraph) is None for paragraph in paragraphs[2:]))
        once = fault.normalize_duplicate_number("1. 1. 重复编号")
        self.assertEqual(once, "1. 重复编号")
        self.assertEqual(fault.normalize_duplicate_number(once), once)
        audit = fault.audit_format_structure(output, fault.validate_input(data))
        self.assertEqual(audit["duplicate_number_paragraphs"], 0)
        self.assertEqual(audit["native_numbered_paragraphs"], 2)

    def test_structured_improvement_plan_and_owner_validation(self) -> None:
        data = fault_data()
        data["improvement_plan"] = [
            {
                "title": "取证增强",
                "prerequisites": ["维护窗口可用"],
                "actions": ["保持 kdump", "配置 netconsole"],
                "owner": "待指定",
                "validation": ["确认 vmcore 及远程日志可落盘"],
                "rollback": [],
            }
        ]
        output = self.directory / "fault-structured-improvement.docx"
        fault.generate(data, fault.DEFAULT_TEMPLATE, output)
        text = all_document_text(output)
        for expected in ("改进项：取证增强", "前提：维护窗口可用", "措施：保持 kdump", "责任人：待指定"):
            self.assertIn(expected, text)

        empty_owner = fault_data()
        empty_owner["improvement_plan"] = [{
            "title": "取证增强",
            "prerequisites": [],
            "actions": ["保持 kdump"],
            "owner": " ",
            "validation": [],
            "rollback": [],
        }]
        with self.assertRaisesRegex(ValueError, "owner"):
            fault.validate_input(empty_owner)

        high_risk = fault_data()
        high_risk["improvement_plan"] = [{
            "title": "高风险配置调整",
            "prerequisites": ["维护窗口可用"],
            "actions": ["调整配置"],
            "owner": "责任团队",
            "validation": ["验证服务状态"],
            "rollback": [],
            "high_risk": True,
        }]
        with self.assertRaisesRegex(ValueError, "rollback"):
            fault.validate_input(high_risk)

    def test_format_audit_failure_does_not_replace_existing_output(self) -> None:
        output = self.directory / "existing.docx"
        original = b"existing output"
        output.write_bytes(original)
        real_audit = fault.audit_format_structure
        try:
            fault.audit_format_structure = lambda path, data: (_ for _ in ()).throw(
                ValueError("forced audit failure")
            )
            with self.assertRaisesRegex(ValueError, "forced audit failure"):
                fault.generate(fault_data(), fault.DEFAULT_TEMPLATE, output)
        finally:
            fault.audit_format_structure = real_audit
        self.assertEqual(output.read_bytes(), original)

    def test_format_audit_rejects_ordinary_all_bold_paragraph(self) -> None:
        data = fault_data()
        output = self.directory / "fault-valid.docx"
        mutated = self.directory / "fault-all-bold.docx"
        fault.generate(data, fault.DEFAULT_TEMPLATE, output)
        document = Document(str(output))
        info_table = fault.unique_table_by_first_cell(document, "故障基本信息")
        _, _, target = fault.description_target(info_table)
        paragraph = next(item for item in target.paragraphs if item.text.strip())
        for run in paragraph.runs:
            if run.text:
                run.bold = True
        document.save(mutated)
        with self.assertRaisesRegex(ValueError, "ordinary paragraph is entirely bold"):
            fault.audit_format_structure(mutated, fault.validate_input(data))

    def test_bom_cli_generation(self) -> None:
        input_path = self.directory / "fault-bom.json"
        output = self.directory / "fault-bom.docx"
        input_path.write_text(json.dumps(fault_data(), ensure_ascii=False), encoding="utf-8-sig")
        result = subprocess.run(
            [sys.executable, str(FAULT_SCRIPT), str(input_path), "--out", str(output)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(output.is_file())
        Document(str(output))

    def test_zip_integrity_after_generation(self) -> None:
        for generator, data, template, name in (
            (fault.generate, fault_data(), fault.DEFAULT_TEMPLATE, "fault.docx"),
            (change.generate, change_data(), change.DEFAULT_TEMPLATE, "change.docx"),
        ):
            with self.subTest(name=name):
                output = self.directory / name
                generator(data, template, output)
                with zipfile.ZipFile(output) as archive:
                    self.assertIsNone(archive.testzip())
                Document(str(output))


class TemplateProbeTests(unittest.TestCase):
    def test_print_template_is_a_real_probe(self) -> None:
        summary = fault.probe_template(Document(str(fault.DEFAULT_TEMPLATE)))
        self.assertEqual(summary["cover_fields"][0]["placeholder"], "云平台数据库残留XXXX")
        self.assertEqual(len(summary["info_fields"]), len(fault.INFO_FIELDS))
        change_summary = change.probe_template(Document(str(change.DEFAULT_TEMPLATE)))
        self.assertEqual(change_summary["checklist"]["row_count"], 65)

    def test_probe_rejects_drifted_template(self) -> None:
        document = Document(str(fault.DEFAULT_TEMPLATE))
        paragraph = next(item for item in document.paragraphs if item.text.strip() == "云平台数据库残留XXXX")
        paragraph.text = "DRIFTED"
        with self.assertRaisesRegex(ValueError, "云平台数据库残留XXXX"):
            fault.probe_template(document)


if __name__ == "__main__":
    unittest.main()
