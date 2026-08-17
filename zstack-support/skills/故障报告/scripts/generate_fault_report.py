#!/usr/bin/env python3
"""Render validated ZStack fault-report JSON into the standard DOCX template."""

from __future__ import annotations

import sys

if sys.version_info < (3, 10):
    print("ERROR: this generator requires Python 3.10 or newer.", file=sys.stderr)
    raise SystemExit(2)

import argparse
import base64
import binascii
from copy import deepcopy
import hashlib
import ipaddress
import json
import os
from pathlib import Path
import posixpath
import re
import tempfile
from typing import Any, Iterable, NamedTuple
import zipfile
from xml.etree import ElementTree as ET
from urllib.parse import unquote, urlsplit


SCRIPT_PATH = Path(__file__).resolve()
SKILL_DIR = SCRIPT_PATH.parents[1]
PLUGIN_DIR = SCRIPT_PATH.parents[3]
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "ZStack企业版-故障分析报告模板V1.2.docx"


def dependency_directories() -> list[Path]:
    override = os.environ.get("ZSTACK_SUPPORT_PYTHONPATH")
    if override:
        return [Path(item).expanduser().resolve() for item in override.split(os.pathsep) if item]
    if sys.platform == "win32":
        base = Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local"))
        return [base / "ZStackSupportAgent" / "python"]
    if sys.platform == "darwin":
        return [Path.home() / "Library" / "Caches" / "ZStackSupportAgent" / "python"]
    return [Path.home() / ".cache" / "ZStackSupportAgent" / "python"]


DEPENDENCY_DIRS = dependency_directories()
for dependency_dir in reversed(DEPENDENCY_DIRS):
    if str(dependency_dir) not in sys.path:
        sys.path.insert(0, str(dependency_dir))

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.table import _Cell
    from docx.text.paragraph import Paragraph
    from lxml import etree as LET
except (ImportError, ModuleNotFoundError) as exc:
    target = DEPENDENCY_DIRS[0]
    requirements = PLUGIN_DIR / "requirements.txt"
    print(
        "ERROR: python-docx is unavailable. Install the plugin's pinned dependencies with:\n"
        f'  "{sys.executable}" -m pip install --target "{target}" -r "{requirements}"\n'
        "Set ZSTACK_SUPPORT_PYTHONPATH to use a different private dependency directory.",
        file=sys.stderr,
    )
    raise SystemExit(2) from exc


INFO_FIELDS = [
    ("project_name", "项目名称"),
    ("customer_name", "客户名称"),
    ("fault_impact_scope", "故障影响和范围"),
    ("customer_contact", "客户联系人"),
    ("reporter", "故障报告人"),
    ("software_product", "软件产品"),
    ("version", "版本号"),
    ("fault_start_time", "故障发生时间"),
    ("business_recovery_time", "业务恢复时间"),
    ("fault_duration", "故障总耗时长"),
    ("fault_category", "故障类别"),
    ("responsible_department", "故障责任部门"),
    ("fault_level", "故障级别"),
]

TEXT_BLOCKS = {
    "analysis_process": "分析处理过程",
    "root_cause_analysis": "原因分析",
    "improvement_plan": "后续改进与预防方案",
}

COVER_FIELDS = {
    "report_title": "云平台数据库残留XXXX",
    "report_subtitle": "——问题故障报告",
}

CONTENT_FIELDS = {"fault_description", *TEXT_BLOCKS}
REQUIRED_FIELDS = {"audience", "report_title", "report_subtitle", *[key for key, _ in INFO_FIELDS], *CONTENT_FIELDS}
ALLOWED_FIELDS = REQUIRED_FIELDS | {"customer_identity_handling"}
AUDIENCES = {"internal", "customer"}
CUSTOMER_IDENTITY_HANDLING = {"confirmed_same_customer", "generalized"}
GENERALIZED_FAULT_IDENTITIES = {
    "project_name": "项目（已泛化）",
    "customer_name": "客户（已泛化）",
    "customer_contact": "客户联系人（已泛化）",
}
IMPROVEMENT_ITEM_FIELDS = {
    "title",
    "prerequisites",
    "actions",
    "owner",
    "validation",
    "rollback",
    "high_risk",
}
REQUIRED_IMPROVEMENT_ITEM_FIELDS = IMPROVEMENT_ITEM_FIELDS - {"high_risk"}
MANUAL_NUMBER_RE = re.compile(r"^\s*\d+\s*[.、](?!\d)\s*")
DUPLICATE_NUMBER_RE = re.compile(r"^\s*(\d+)\s*[.、](?!\d)\s*\1\s*[.、](?!\d)\s*")
LABEL_RE = re.compile(r"^[^:：\r\n]{1,24}[:：]")
TEMPLATE_PLACEHOLDERS = {COVER_FIELDS["report_title"]}


class RenderLine(NamedTuple):
    text: str
    numbered: bool = False
    line_bold: bool = False

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CORE_NS = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
DC_NS = "http://purl.org/dc/elements/1.1/"
DCTERMS_NS = "http://purl.org/dc/terms/"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL_BASE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
HYPERLINK_REL_TYPE = f"{OFFICE_REL_BASE}/hyperlink".casefold()
ALLOWED_INTERNAL_REL_TYPES = {
    "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties",
    *(f"{OFFICE_REL_BASE}/{role}" for role in (
        "officeDocument", "fontTable", "footer", "header", "image", "numbering", "settings",
        "styles", "theme", "endnotes", "footnotes", "webSettings",
    )),
}
ALLOWED_INTERNAL_REL_TYPES = {value.casefold() for value in ALLOWED_INTERNAL_REL_TYPES}
URL_RE = re.compile(r"https?://[^\s<>()\[\]{}\"']+", re.IGNORECASE)
DOTTED_IPV4_RE = re.compile(
    r"(?<![A-Za-z0-9_.])(?:0x[0-9a-f]+|\d+)(?:\.(?:0x[0-9a-f]+|\d+)){1,3}(?![A-Za-z0-9_.])",
    re.IGNORECASE,
)
INTEGER_IPV4_RE = re.compile(r"(?<![A-Za-z0-9_.])(?:0x[0-9a-f]{1,8}|0[0-7]{7,11}|\d{8,10})(?![A-Za-z0-9_.])", re.IGNORECASE)
IPV6_RE = re.compile(r"(?<![0-9a-f:])(?:[0-9a-f]{0,4}:){2,7}[0-9a-f]{0,4}(?![0-9a-f:])", re.IGNORECASE)
UUID_RE = re.compile(r"(?<![0-9a-f])[0-9a-f]{8}(?:-[0-9a-f]{4}){3}-[0-9a-f]{12}(?![0-9a-f])", re.IGNORECASE)
COMPACT_GUID_RE = re.compile(r"(?<![0-9a-f])[0-9a-f]{12}[1-8][0-9a-f]{3}[89ab][0-9a-f]{15}(?![0-9a-f])", re.IGNORECASE)
INTERNAL_TICKET_RE = re.compile(r"(?<![A-Za-z0-9_])(?:TIC|BUG|SUG|JIRA)-\d+(?![A-Za-z0-9_])", re.IGNORECASE)
RELATIVE_INTERNAL_LINK_RE = re.compile(
    r"(?:forum\.php\?[^\s]*\btid=|/(?:browse|issues?)/|(?:^|[/])(?:bbs|jira|confluence)(?:[/])|confluence/(?:pages|display)/)",
    re.IGNORECASE,
)
INTERNAL_HOST_RE = re.compile(
    r"(?<![A-Za-z0-9_.-])(?:localhost|[A-Za-z0-9_-]+(?:\.[A-Za-z0-9_-]+)*\.(?:internal|local|corp)|(?:bbs|jira|confluence)(?:[.-][A-Za-z0-9_-]+)+)(?::\d+)?",
    re.IGNORECASE,
)
PEM_PRIVATE_KEY_RE = re.compile(
    r"-----BEGIN (?:(?:RSA|EC|DSA|OPENSSH|ENCRYPTED) )?PRIVATE KEY-----|"
    r"-----BEGIN PGP PRIVATE KEY BLOCK-----",
    re.IGNORECASE,
)
AUTHORIZATION_SECRET_RE = re.compile(
    r"(?<![A-Za-z0-9_])Authorization\s*[:=：]\s*(?:Bearer|Basic)\s+[^\s,;]+",
    re.IGNORECASE,
)
BASIC_CREDENTIAL_RE = re.compile(r"(?<![A-Za-z0-9_])Basic\s+([A-Za-z0-9+/]+={0,2})(?![A-Za-z0-9+/=])", re.IGNORECASE)
SECRET_ASSIGNMENT_RE = re.compile(
    r"(?<![A-Za-z0-9_])(?:password|passwd|pwd|client[_-]?secret|secret|api[_-]?key|access[_-]?key|auth[_-]?token|token)(?![A-Za-z0-9_])\s*[:=：]\s*[\"']?([^\s\"',;；]+)",
    re.IGNORECASE,
)
URI_USERINFO_RE = re.compile(r"https?://[^/\s:@]+:[^@\s/]+@", re.IGNORECASE)
KNOWN_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9_])(?:AKIA[0-9A-Z]{16}|ghp_[A-Za-z0-9]{20,}|github_pat_[A-Za-z0-9_]{20,})(?![A-Za-z0-9_])")
JWT_RE = re.compile(r"(?<![A-Za-z0-9_])eyJ[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{8,}(?![A-Za-z0-9_])")


def plain_text(value: str) -> str:
    return value.replace("**", "").strip()


def validate_nonempty_string(value: Any, field: str) -> str:
    if not isinstance(value, str):
        raise ValueError(f"Field '{field}' must be a string.")
    value = value.strip()
    if not value or not plain_text(value):
        raise ValueError(f"Field '{field}' must not be empty; use an explicit value such as '待补充'.")
    return value


def validate_string_array(value: Any, field: str, *, allow_empty: bool = False) -> list[str]:
    if not isinstance(value, list):
        raise ValueError(f"Field '{field}' must be an array of strings.")
    if not value and not allow_empty:
        raise ValueError(f"Field '{field}' must not be an empty array.")
    return [validate_nonempty_string(item, f"{field}[{index}]") for index, item in enumerate(value)]


def validate_improvement_item(value: Any, field: str) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise ValueError(f"Field '{field}' must be a string or a structured improvement object.")
    unknown = sorted(set(value) - IMPROVEMENT_ITEM_FIELDS)
    missing = sorted(REQUIRED_IMPROVEMENT_ITEM_FIELDS - set(value))
    if unknown:
        raise ValueError(f"Unknown field(s) in '{field}': {', '.join(unknown)}.")
    if missing:
        raise ValueError(
            f"Missing required field(s) in '{field}': {', '.join(missing)}. "
            "Use owner='待指定' when ownership is not confirmed."
        )
    validated = {
        "title": validate_nonempty_string(value["title"], f"{field}.title"),
        "prerequisites": validate_string_array(
            value["prerequisites"], f"{field}.prerequisites", allow_empty=True
        ),
        "actions": validate_string_array(value["actions"], f"{field}.actions"),
        "owner": validate_nonempty_string(value["owner"], f"{field}.owner"),
        "validation": validate_string_array(
            value["validation"], f"{field}.validation", allow_empty=True
        ),
        "rollback": validate_string_array(value["rollback"], f"{field}.rollback", allow_empty=True),
    }
    high_risk = value.get("high_risk", False)
    if not isinstance(high_risk, bool):
        raise ValueError(f"Field '{field}.high_risk' must be a JSON boolean.")
    validated["high_risk"] = high_risk
    if high_risk:
        missing_controls = [
            name
            for name in ("prerequisites", "validation", "rollback")
            if not validated[name]
        ]
        if missing_controls:
            raise ValueError(
                f"High-risk improvement '{field}' requires non-empty: {', '.join(missing_controls)}."
            )
    return validated


def validate_content(value: Any, field: str) -> str | list[str] | list[str | dict[str, Any]]:
    if isinstance(value, str):
        validated = validate_nonempty_string(value, field)
        if not [line for line in validated.splitlines() if line.strip()]:
            raise ValueError(f"Field '{field}' must contain text.")
        return validated
    if not isinstance(value, list):
        raise ValueError(f"Field '{field}' must be a string or a non-empty array of strings.")
    if not value:
        raise ValueError(f"Field '{field}' must not be an empty array.")
    if field == "improvement_plan":
        return [
            validate_nonempty_string(item, f"{field}[{index}]")
            if isinstance(item, str)
            else validate_improvement_item(item, f"{field}[{index}]")
            for index, item in enumerate(value)
        ]
    return [validate_nonempty_string(item, f"{field}[{index}]") for index, item in enumerate(value)]


def iter_text_values(value: Any, path: str = "$") -> Iterable[tuple[str, str]]:
    if isinstance(value, str):
        yield path, value
    elif isinstance(value, dict):
        for key, item in value.items():
            yield from iter_text_values(item, f"{path}.{key}")
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from iter_text_values(item, f"{path}[{index}]")


def decoded_text(text: str) -> str:
    values = [text]
    current = text
    for _ in range(3):
        decoded = unquote(current)
        if decoded == current:
            break
        values.append(decoded)
        current = decoded
    return "\n".join(values)


def validate_no_secrets(text: str, field: str) -> None:
    scan = decoded_text(text)
    if PEM_PRIVATE_KEY_RE.search(scan):
        raise ValueError(f"Output field '{field}' contains private-key PEM material.")
    if AUTHORIZATION_SECRET_RE.search(scan) or URI_USERINFO_RE.search(scan):
        raise ValueError(f"Output field '{field}' contains an authorization credential.")
    for match in BASIC_CREDENTIAL_RE.finditer(scan):
        try:
            decoded = base64.b64decode(match.group(1), validate=True)
        except (binascii.Error, ValueError):
            continue
        if b":" in decoded:
            raise ValueError(f"Output field '{field}' contains an authorization credential.")
    if KNOWN_TOKEN_RE.search(scan) or JWT_RE.search(scan):
        raise ValueError(f"Output field '{field}' contains an access token.")
    safe_values = {
        "待补充", "待确认", "已脱敏", "redacted", "<redacted>", "***", "none", "null",
        "disabled", "unset", "unavailable", "omitted", "false", "off", "not-set", "not_configured",
    }
    for match in SECRET_ASSIGNMENT_RE.finditer(scan):
        if match.group(1).strip().casefold() not in {value.casefold() for value in safe_values}:
            raise ValueError(f"Output field '{field}' contains an assigned password, token, or secret.")


def parse_ipv4_number(value: str) -> int | None:
    try:
        if value.casefold().startswith("0x"):
            number = int(value, 16)
        elif len(value) > 1 and value.startswith("0"):
            number = int(value, 8) if all(character in "01234567" for character in value) else int(value, 10)
        else:
            number = int(value, 10)
    except ValueError:
        return None
    return number if 0 <= number <= 0xFFFFFFFF else None


def parse_legacy_ipv4(value: str) -> int | None:
    components = value.split(".")
    if not 1 <= len(components) <= 4 or any(not component for component in components):
        return None
    numbers = [parse_ipv4_number(component) for component in components]
    if any(number is None for number in numbers):
        return None
    parsed = [int(number) for number in numbers if number is not None]
    limits = {
        1: (0xFFFFFFFF,),
        2: (0xFF, 0xFFFFFF),
        3: (0xFF, 0xFF, 0xFFFF),
        4: (0xFF, 0xFF, 0xFF, 0xFF),
    }[len(parsed)]
    if any(number > limit for number, limit in zip(parsed, limits)):
        return None
    result = parsed[-1]
    shift = 8 * (4 - len(parsed) + 1)
    for number in reversed(parsed[:-1]):
        result |= number << shift
        shift += 8
    return result


def version_context(text: str, start: int) -> bool:
    prefix = text[max(0, start - 32):start]
    return bool(re.search(r"(?i)(?:version|ver|release|build|版本|构建)\s*[:=v-]?\s*$", prefix))


def semantic_version_candidate(value: str) -> bool:
    parts = value.split(".")
    return (
        len(parts) == 3
        and all(part.isdigit() and (part == "0" or not part.startswith("0")) for part in parts)
        and int(parts[0]) <= 20
        and all(int(part) <= 99 for part in parts[1:])
    )


def ip_number_context(text: str, start: int, end: int) -> bool:
    prefix = text[max(0, start - 32):start]
    suffix = text[end:min(len(text), end + 12)]
    return bool(
        re.search(
            r"(?i)(?:(?<![A-Za-z0-9_])(?:ip(?:v4)?|host|endpoint)(?![A-Za-z0-9_])|地址|端点)"
            r"\s*[:=：]?\s*$",
            prefix,
        )
        or re.match(r"^:\d{1,5}\b", suffix)
    )


def validate_customer_text(text: str, field: str) -> None:
    text = decoded_text(text)
    lowered = text.casefold()
    if UUID_RE.search(text) or COMPACT_GUID_RE.search(text):
        raise ValueError(f"Customer output field '{field}' contains a raw GUID; generalize the identifier.")
    if INTERNAL_TICKET_RE.search(text):
        raise ValueError(f"Customer output field '{field}' contains an internal ticket identifier.")
    if RELATIVE_INTERNAL_LINK_RE.search(URL_RE.sub("", text)):
        raise ValueError(f"Customer output field '{field}' contains a relative internal-system link.")
    if "file://" in lowered or re.search(r"\\\\[^\\\s]+\\[^\s]+", text) or re.search(r"\b[A-Za-z]:\\", text):
        raise ValueError(f"Customer output field '{field}' contains an internal file endpoint.")
    if INTERNAL_HOST_RE.search(text):
        raise ValueError(f"Customer output field '{field}' contains an internal BBS/Jira/Confluence or private host.")
    for match in DOTTED_IPV4_RE.finditer(text):
        candidate = match.group(0)
        if (
            parse_legacy_ipv4(candidate) is None
            or version_context(text, match.start())
            or semantic_version_candidate(candidate)
            or field.casefold().endswith(".version")
        ):
            continue
        if candidate.count(".") == 1 and not (
            candidate.casefold().startswith("0x")
            or candidate.startswith("0")
            or ip_number_context(text, match.start(), match.end())
        ):
            continue
        raise ValueError(f"Customer output field '{field}' contains a raw IP address; generalize the endpoint.")
    for match in INTEGER_IPV4_RE.finditer(text):
        candidate = match.group(0)
        encoded_ip_form = (
            candidate.casefold().startswith("0x") and len(candidate[2:]) == 8
        ) or (
            candidate.startswith("0") and len(candidate) >= 8 and all(character in "01234567" for character in candidate)
        )
        if parse_ipv4_number(candidate) is not None and (
            encoded_ip_form or ip_number_context(text, match.start(), match.end())
        ):
            raise ValueError(f"Customer output field '{field}' contains an integer-form IPv4 address.")
    for match in IPV6_RE.finditer(text):
        candidate = match.group(0)
        try:
            ipaddress.ip_address(candidate)
        except ValueError:
            continue
        raise ValueError(f"Customer output field '{field}' contains a raw IPv6 address; generalize the endpoint.")
    for match in URL_RE.finditer(text):
        url = match.group(0).rstrip(".,;:!?)]}，。；：！？")
        parsed = urlsplit(url)
        host = (parsed.hostname or "").casefold()
        target = f"{host}{parsed.path}".casefold()
        if any(token in target for token in ("bbs", "jira", "confluence", "atlassian.net")):
            raise ValueError(f"Customer output field '{field}' contains an internal support-system link.")
        try:
            address = ipaddress.ip_address(host)
        except ValueError:
            address = None
        if address is not None or parse_legacy_ipv4(host) is not None:
            raise ValueError(f"Customer output field '{field}' contains a raw IP URL endpoint.")


def validate_customer_content(data: dict[str, Any]) -> None:
    for field, text in iter_text_values(data):
        validate_customer_text(text, field)


def validate_secret_content(data: dict[str, Any]) -> None:
    for field, text in iter_text_values(data):
        validate_no_secrets(text, field)


def validate_customer_xml_text(root: ET.Element, field: str) -> None:
    text_nodes = [
        node.text or ""
        for node in root.iter()
        if node.tag.rsplit("}", 1)[-1].casefold() in {"t", "instrtext"} and (node.text or "")
    ]
    for index, value in enumerate(text_nodes):
        validate_customer_text(value, f"{field}:text[{index}]")
    if text_nodes:
        validate_customer_text("\n".join(text_nodes), field)
    for index, paragraph in enumerate(
        node for node in root.iter() if node.tag.rsplit("}", 1)[-1].casefold() == "p"
    ):
        rendered = "".join(
            node.text or ""
            for node in paragraph.iter()
            if node.tag.rsplit("}", 1)[-1].casefold() in {"t", "instrtext"}
        )
        if rendered:
            validate_customer_text(rendered, f"{field}:paragraph[{index}]")


def validate_secret_xml(root: ET.Element, field: str) -> None:
    values = [node.text or "" for node in root.iter() if node.text]
    values.extend(value for node in root.iter() for value in node.attrib.values() if value)
    for index, value in enumerate(values):
        validate_no_secrets(value, f"{field}:value[{index}]")
    if values:
        validate_no_secrets("\n".join(values), field)
        validate_no_secrets("".join(values), field)
    for index, paragraph in enumerate(
        node for node in root.iter() if node.tag.rsplit("}", 1)[-1].casefold() == "p"
    ):
        rendered = "".join(node.text or "" for node in paragraph.iter() if node.text)
        if rendered:
            validate_no_secrets(rendered, f"{field}:paragraph[{index}]")


def validate_input(data: Any) -> dict[str, Any]:
    if not isinstance(data, dict):
        raise ValueError("Input JSON must be an object.")
    unknown = sorted(set(data) - ALLOWED_FIELDS)
    missing = sorted(REQUIRED_FIELDS - set(data))
    if unknown:
        raise ValueError(f"Unknown input field(s): {', '.join(unknown)}.")
    if missing:
        raise ValueError(f"Missing required field(s): {', '.join(missing)}.")

    validated: dict[str, Any] = {}
    audience = validate_nonempty_string(data["audience"], "audience")
    if audience not in AUDIENCES:
        raise ValueError("Field 'audience' must be 'internal' or 'customer'.")
    validated["audience"] = audience
    for field in sorted(REQUIRED_FIELDS - CONTENT_FIELDS - {"audience"}):
        validated[field] = validate_nonempty_string(data[field], field)
    for field in sorted(CONTENT_FIELDS):
        validated[field] = validate_content(data[field], field)
    root_cause_text = plain_text(
        "\n".join(validated["root_cause_analysis"])
        if isinstance(validated["root_cause_analysis"], list)
        else validated["root_cause_analysis"]
    )
    evidence_markers = (
        "证据边界",
        "结论状态",
        "证据支持",
        "初步判断",
        "仍需确认",
        "待确认",
        "未闭环",
        "已闭环",
        "已确认",
    )
    if not any(marker in root_cause_text for marker in evidence_markers):
        raise ValueError(
            "Field 'root_cause_analysis' must state an evidence boundary or conclusion status "
            "(for example: '证据边界：待补充' or '初步判断，仍需确认')."
        )
    if audience == "customer":
        if "customer_identity_handling" not in data:
            raise ValueError(
                "Customer output requires 'customer_identity_handling' to be "
                "'confirmed_same_customer' or 'generalized'."
            )
        identity_handling = validate_nonempty_string(
            data["customer_identity_handling"], "customer_identity_handling"
        )
        if identity_handling not in CUSTOMER_IDENTITY_HANDLING:
            raise ValueError(
                "Field 'customer_identity_handling' must be 'confirmed_same_customer' or 'generalized'."
            )
        validated["customer_identity_handling"] = identity_handling
        if identity_handling == "generalized":
            for field, required_value in GENERALIZED_FAULT_IDENTITIES.items():
                if validated[field] != required_value:
                    raise ValueError(
                        f"Customer generalized output requires '{field}' to be exactly '{required_value}'."
                    )
        validate_customer_content(validated)
    elif "customer_identity_handling" in data:
        raise ValueError("Field 'customer_identity_handling' is only valid when audience is 'customer'.")
    validate_secret_content(validated)
    return validated


def normalize_duplicate_number(text: str) -> str:
    pattern = re.compile(
        r"^(\s*(?:\*\*)?(\d+)\s*[.、](?!\d)(?:\*\*)?\s*)"
        r"(?:\*\*)?\2\s*[.、](?!\d)(?:\*\*)?\s*"
    )
    normalized = text
    while True:
        match = pattern.match(normalized)
        if match is None:
            return normalized
        normalized = match.group(1) + normalized[match.end():]


def string_render_lines(value: str | list[str]) -> list[RenderLine]:
    source = value if isinstance(value, list) else [line.strip() for line in value.splitlines() if line.strip()]
    lines: list[RenderLine] = []
    for item in source:
        text = normalize_duplicate_number(item)
        normalized = plain_text(text)
        manually_numbered = bool(MANUAL_NUMBER_RE.match(normalized))
        label_candidate = MANUAL_NUMBER_RE.sub("", normalized, count=1)
        has_label = bool(LABEL_RE.match(label_candidate))
        lines.append(
            RenderLine(
                text=text,
                numbered=not manually_numbered and not has_label,
                line_bold=bool(re.fullmatch(r"\*\*.+\*\*", text.strip(), flags=re.DOTALL)),
            )
        )
    return lines


def improvement_render_lines(value: str | list[str | dict[str, Any]]) -> list[RenderLine]:
    if isinstance(value, str):
        return string_render_lines(value)
    lines: list[RenderLine] = []
    for item in value:
        if isinstance(item, str):
            lines.extend(string_render_lines([item]))
            continue
        lines.append(RenderLine(f"改进项：{item['title']}", line_bold=True))
        for label, key in (
            ("前提", "prerequisites"),
            ("措施", "actions"),
            ("验证标准", "validation"),
            ("回退条件", "rollback"),
        ):
            for entry in item[key]:
                lines.append(RenderLine(f"{label}：{entry}"))
        lines.append(RenderLine(f"责任人：{item['owner']}"))
    return lines


def format_items(
    value: str | list[str] | list[str | dict[str, Any]],
    field: str,
) -> list[RenderLine]:
    if field == "improvement_plan":
        return improvement_render_lines(value)
    return string_render_lines(value)  # type: ignore[arg-type]


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def run_prototype(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    ppr = paragraph._p.pPr
    paragraph_rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
    if paragraph_rpr is not None:
        return deepcopy(paragraph_rpr)
    return None


def add_run(paragraph: Paragraph, text: str, prototype: Any, bold: bool = False) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    if prototype is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, deepcopy(prototype))
    run.bold = bool(bold)


def write_markdown_segments(
    paragraph: Paragraph,
    text: str,
    prototype: Any,
    *,
    line_bold: bool,
) -> None:
    for segment in re.split(r"(\*\*.*?\*\*)", text):
        if not segment:
            continue
        marked = segment.startswith("**") and segment.endswith("**") and len(segment) >= 4
        rendered = segment[2:-2] if marked else segment
        add_run(paragraph, rendered, prototype, bold=line_bold or marked)


def write_rich_text(
    paragraph: Paragraph,
    text: str,
    line_bold: bool = False,
    prototype: Any = None,
    label_bold: bool = True,
) -> None:
    prototype = deepcopy(prototype) if prototype is not None else run_prototype(paragraph)
    clear_paragraph(paragraph)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    number_match = re.match(
        r"^(\s*(?:\*\*)?\d+\s*[.、](?!\d)(?:\*\*)?\s*)(.*)$",
        text,
        flags=re.DOTALL,
    )
    prefix = number_match.group(1) if number_match else ""
    body = number_match.group(2) if number_match else text
    if prefix:
        write_markdown_segments(paragraph, prefix, prototype, line_bold=line_bold)

    label_match = re.match(r"^([^:*：\r\n]{1,24}[:：])(.*)$", body, flags=re.DOTALL)
    if label_bold and label_match:
        add_run(paragraph, label_match.group(1), prototype, bold=True)
        write_markdown_segments(
            paragraph,
            label_match.group(2),
            prototype,
            line_bold=line_bold,
        )
    else:
        write_markdown_segments(paragraph, body, prototype, line_bold=line_bold)


def paragraph_prototype(paragraph: Paragraph) -> tuple[Any, Any]:
    return (
        deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None,
        run_prototype(paragraph),
    )


def apply_paragraph_prototype(paragraph: Paragraph, prototype: tuple[Any, Any]) -> None:
    ppr, _ = prototype
    if paragraph._p.pPr is not None:
        paragraph._p.remove(paragraph._p.pPr)
    if ppr is not None:
        paragraph._p.insert(0, deepcopy(ppr))


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def decimal_numbering_abstract_id(doc: Any) -> int:
    numbering = doc.part.numbering_part.element
    for abstract in numbering.findall(qn("w:abstractNum")):
        level = next(
            (item for item in abstract.findall(qn("w:lvl")) if item.get(qn("w:ilvl")) == "0"),
            None,
        )
        if level is None:
            continue
        number_format = level.find(qn("w:numFmt"))
        level_text = level.find(qn("w:lvlText"))
        if (
            number_format is not None
            and number_format.get(qn("w:val")) == "decimal"
            and level_text is not None
            and "%1" in level_text.get(qn("w:val"), "")
        ):
            return int(abstract.get(qn("w:abstractNumId")))
    raise ValueError("Template does not contain a reusable decimal numbering definition.")


def create_numbering_instance(doc: Any) -> int:
    numbering = doc.part.numbering_part.element
    existing = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    number_id = max(existing, default=0) + 1
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(number_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(decimal_numbering_abstract_id(doc)))
    number.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    number.append(override)
    numbering.append(number)
    return number_id


def apply_native_number(paragraph: Paragraph, number_id: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    existing = properties.find(qn("w:numPr"))
    if existing is not None:
        properties.remove(existing)
    number_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(number_id))
    number_properties.append(level)
    number_properties.append(number)
    properties.append(number_properties)


def set_cell_lines(doc: Any, cell: _Cell, lines: list[RenderLine], line_bold: bool = False) -> None:
    paragraphs = list(cell.paragraphs)
    prototype = paragraph_prototype(paragraphs[0])
    for paragraph in paragraphs[1:]:
        remove_paragraph(paragraph)
    number_id = create_numbering_instance(doc) if any(line.numbered for line in lines) else None
    apply_paragraph_prototype(paragraphs[0], prototype)
    write_rich_text(
        paragraphs[0],
        lines[0].text,
        line_bold=line_bold or lines[0].line_bold,
        prototype=prototype[1],
    )
    if lines[0].numbered and number_id is not None:
        apply_native_number(paragraphs[0], number_id)
    for line in lines[1:]:
        paragraph = cell.add_paragraph()
        apply_paragraph_prototype(paragraph, prototype)
        write_rich_text(paragraph, line.text, line_bold=line.line_bold, prototype=prototype[1])
        if line.numbered and number_id is not None:
            apply_native_number(paragraph, number_id)


def set_cell_text(doc: Any, cell: _Cell, text: str, line_bold: bool = False) -> None:
    set_cell_lines(doc, cell, [RenderLine(text, line_bold=line_bold)], line_bold=line_bold)


def unique_table_by_first_cell(doc: Any, first_cell_text: str):
    matches = [
        table
        for table in doc.tables
        if table.rows and table.rows[0].cells and table.rows[0].cells[0].text.strip() == first_cell_text
    ]
    if len(matches) != 1:
        raise ValueError(f"Template must contain exactly one table titled '{first_cell_text}', found {len(matches)}.")
    return matches[0]


def unique_paragraph(doc: Any, text: str) -> tuple[int, Paragraph]:
    matches = [(index, paragraph) for index, paragraph in enumerate(doc.paragraphs) if paragraph.text.strip() == text]
    if len(matches) != 1:
        raise ValueError(f"Template must contain exactly one paragraph '{text}', found {len(matches)}.")
    return matches[0]


def label_target(table: Any, label: str) -> tuple[int, int, _Cell]:
    candidates: list[tuple[int, int, _Cell]] = []
    seen: set[tuple[int, int]] = set()
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            if cell.text.strip() != label or column_index + 1 >= len(row.cells):
                continue
            target = row.cells[column_index + 1]
            if target._tc is cell._tc:
                continue
            key = (row_index, id(target._tc))
            if key not in seen:
                seen.add(key)
                candidates.append((row_index, column_index + 1, target))
    if len(candidates) > 1:
        single_column_targets = [
            candidate
            for candidate in candidates
            if sum(cell._tc is candidate[2]._tc for cell in table.rows[candidate[0]].cells) == 1
        ]
        if len(single_column_targets) == 1:
            candidates = single_column_targets
    if len(candidates) > 1:
        blank_targets = [candidate for candidate in candidates if not candidate[2].text.strip()]
        if len(blank_targets) == 1:
            candidates = blank_targets
    if len(candidates) != 1:
        raise ValueError(f"Template label '{label}' has {len(candidates)} writable target candidates.")
    return candidates[0]


def description_target(info_table: Any) -> tuple[int, int, _Cell]:
    matching_rows = [
        index
        for index, row in enumerate(info_table.rows)
        if row.cells and row.cells[0].text.strip() == "故障情况描述"
    ]
    if len(matching_rows) != 1:
        raise ValueError(f"Template must contain one '故障情况描述' row, found {len(matching_rows)}.")
    target_row = matching_rows[0] + 1
    if target_row >= len(info_table.rows) or not info_table.rows[target_row].cells:
        raise ValueError("Template fault-description row has no writable content row.")
    return target_row, 0, info_table.rows[target_row].cells[0]


def probe_template(doc: Any) -> dict[str, Any]:
    covers = []
    for key, placeholder in COVER_FIELDS.items():
        index, paragraph = unique_paragraph(doc, placeholder)
        covers.append({"key": key, "placeholder": paragraph.text, "paragraph_index": index})

    info_table = unique_table_by_first_cell(doc, "故障基本信息")
    info = []
    for key, label in INFO_FIELDS:
        row, column, target = label_target(info_table, label)
        info.append(
            {
                "key": key,
                "label": label,
                "target_row": row,
                "target_column": column,
                "current_value": target.text.strip(),
            }
        )
    row, column, target = description_target(info_table)

    blocks = []
    for key, title in TEXT_BLOCKS.items():
        table = unique_table_by_first_cell(doc, title)
        if len(table.rows) < 2 or not table.rows[1].cells:
            raise ValueError(f"Template text block '{title}' has no writable content row.")
        blocks.append({"key": key, "title": title, "current_value": table.rows[1].cells[0].text.strip()})

    return {
        "cover_fields": covers,
        "info_fields": info,
        "fault_description": {
            "key": "fault_description",
            "target_row": row,
            "target_column": column,
            "current_value": target.text.strip(),
        },
        "text_blocks": blocks,
    }


def update_document(doc: Any, data: dict[str, Any]) -> None:
    probe_template(doc)
    for key, placeholder in COVER_FIELDS.items():
        _, paragraph = unique_paragraph(doc, placeholder)
        write_rich_text(paragraph, data[key], line_bold=key == "report_title")

    info_table = unique_table_by_first_cell(doc, "故障基本信息")
    for key, label in INFO_FIELDS:
        _, _, target = label_target(info_table, label)
        set_cell_text(
            doc,
            target,
            data[key],
            line_bold=key in {"fault_level", "fault_start_time", "business_recovery_time"},
        )
    _, _, target = description_target(info_table)
    set_cell_lines(doc, target, format_items(data["fault_description"], "fault_description"))

    for key, title in TEXT_BLOCKS.items():
        table = unique_table_by_first_cell(doc, title)
        set_cell_lines(doc, table.rows[1].cells[0], format_items(data[key], key))


def iter_cell_paragraphs(cell: _Cell) -> Iterable[Paragraph]:
    for paragraph in cell.paragraphs:
        yield paragraph
    for table in cell.tables:
        for row in table.rows:
            for nested_cell in row.cells:
                yield from iter_cell_paragraphs(nested_cell)


def iter_all_paragraphs(doc: Any) -> Iterable[Paragraph]:
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_cell_paragraphs(cell)


def apply_line_spacing(doc: Any) -> None:
    for paragraph in iter_all_paragraphs(doc):
        paragraph.paragraph_format.line_spacing = 1.5


def mark_fields_for_update(doc: Any) -> None:
    update_fields = doc.settings.element.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        doc.settings.element.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    for field in doc.element.iter(qn("w:fldChar")):
        if field.get(qn("w:fldCharType")) == "begin":
            field.set(qn("w:dirty"), "true")


def clean_core_properties(doc: Any, title: str, subject: str) -> None:
    props = doc.core_properties
    props.title = plain_text(title)
    props.subject = subject
    props.author = ""
    props.last_modified_by = ""
    props.revision = 1


def sanitize_xml_part(name: str, data: bytes) -> bytes:
    if not (name.endswith(".xml") or name.endswith(".rels")):
        return data
    root = LET.fromstring(data)
    if name == "docProps/core.xml":
        sensitive_tags = {
            f"{{{DC_NS}}}creator",
            f"{{{CORE_NS}}}lastModifiedBy",
            f"{{{CORE_NS}}}revision",
            f"{{{DCTERMS_NS}}}created",
            f"{{{DCTERMS_NS}}}modified",
        }
        for child in list(root):
            if child.tag in sensitive_tags:
                root.remove(child)
    elif name.endswith(".rels"):
        for relationship in list(root):
            target = relationship.get("Target", "").replace("\\", "/").lower()
            rel_type = relationship.get("Type", "").lower()
            if (
                "customxml" in target
                or target.endswith("docprops/custom.xml")
                or target.endswith("docprops/app.xml")
                or "docprops/thumbnail." in target
                or "custom-properties" in rel_type
                or "extended-properties" in rel_type
                or rel_type.endswith("/metadata/thumbnail")
            ):
                root.remove(relationship)
    elif name == "[Content_Types].xml":
        for child in list(root):
            part_name = child.get("PartName", "").lower()
            if (
                part_name.startswith("/customxml/")
                or part_name.startswith("/docprops/thumbnail.")
                or part_name in {"/docprops/custom.xml", "/docprops/app.xml"}
            ):
                root.remove(child)
    for element in list(root.iter()):
        for attribute in list(element.attrib):
            local_name = LET.QName(attribute).localname.casefold()
            if local_name.startswith("rsid") or local_name in {"durableid", "paraid", "textid"}:
                del element.attrib[attribute]
        if name == "word/settings.xml" and element is not root:
            local_name = LET.QName(element).localname.casefold()
            if local_name in {"docvars", "rsids", "docid"}:
                element.getparent().remove(element)
    if name == "word/settings.xml":
        LET.cleanup_namespaces(root)
        ignorable = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"
        if ignorable in root.attrib:
            declared = {prefix for prefix in root.nsmap if prefix}
            tokens = [token for token in root.get(ignorable, "").split() if token in declared]
            if tokens:
                root.set(ignorable, " ".join(tokens))
            else:
                del root.attrib[ignorable]
    return LET.tostring(root, encoding="utf-8", xml_declaration=True, standalone=True)


def sanitize_package(source: Path, destination: Path) -> None:
    with zipfile.ZipFile(source, "r") as archive_in, zipfile.ZipFile(destination, "w") as archive_out:
        for item in archive_in.infolist():
            normalized = item.filename.replace("\\", "/").lower()
            if (
                normalized in {"docprops/custom.xml", "docprops/app.xml"}
                or normalized.startswith("customxml/")
                or normalized.startswith("docprops/thumbnail.")
            ):
                continue
            archive_out.writestr(item, sanitize_xml_part(item.filename, archive_in.read(item.filename)))


def package_content_types(archive: zipfile.ZipFile) -> tuple[dict[str, str], dict[str, str]]:
    root = ET.fromstring(archive.read("[Content_Types].xml"))
    defaults: dict[str, str] = {}
    overrides: dict[str, str] = {}
    for element in root:
        local_name = element.tag.rsplit("}", 1)[-1]
        if local_name == "Default":
            defaults[element.get("Extension", "").casefold()] = element.get("ContentType", "")
        elif local_name == "Override":
            overrides[element.get("PartName", "").lstrip("/")] = element.get("ContentType", "")
    return defaults, overrides


def part_content_type(part: str, defaults: dict[str, str], overrides: dict[str, str]) -> str:
    if part in overrides:
        return overrides[part]
    extension = part.rsplit(".", 1)[-1].casefold() if "." in part else ""
    return defaults.get(extension, "")


def relationship_part_name(source_part: str | None) -> str:
    if source_part is None:
        return "_rels/.rels"
    directory, filename = posixpath.split(source_part)
    return posixpath.join(directory, "_rels", f"{filename}.rels")


def resolve_relationship_target(source_part: str | None, target: str) -> str:
    decoded = unquote(target).split("#", 1)[0].split("?", 1)[0].replace("\\", "/")
    if decoded.startswith("/"):
        resolved = posixpath.normpath(decoded.lstrip("/"))
    else:
        resolved = posixpath.normpath(posixpath.join(posixpath.dirname(source_part or ""), decoded))
    if resolved == ".." or resolved.startswith("../"):
        raise ValueError(f"DOCX relationship escapes the package root: {target}.")
    return resolved


def read_relationships(archive: zipfile.ZipFile, source_part: str | None) -> list[dict[str, Any]]:
    relationship_part = relationship_part_name(source_part)
    if relationship_part not in archive.namelist():
        return []
    root = ET.fromstring(archive.read(relationship_part))
    relationships = []
    for element in root:
        relationships.append(
            {
                "source": source_part,
                "relationship_part": relationship_part,
                "type": element.get("Type", ""),
                "role": element.get("Type", "").rsplit("/", 1)[-1].casefold(),
                "target": element.get("Target", ""),
                "external": element.get("TargetMode", "").casefold() == "external",
            }
        )
    return relationships


def package_relationship_graph(
    archive: zipfile.ZipFile,
) -> tuple[set[str], dict[str, set[str]], list[dict[str, Any]]]:
    names = set(archive.namelist())
    reachable: set[str] = set()
    roles: dict[str, set[str]] = {}
    external: list[dict[str, Any]] = []
    queue: list[str] = []
    for relationship in read_relationships(archive, None):
        if relationship["external"]:
            external.append(relationship)
            continue
        target = resolve_relationship_target(None, relationship["target"])
        if target not in names:
            raise ValueError(f"DOCX root relationship references missing part '{target}'.")
        roles.setdefault(target, set()).add(relationship["role"])
        queue.append(target)
    while queue:
        part = queue.pop()
        if part in reachable:
            continue
        reachable.add(part)
        for relationship in read_relationships(archive, part):
            if relationship["external"]:
                external.append(relationship)
                continue
            target = resolve_relationship_target(part, relationship["target"])
            if target not in names:
                raise ValueError(f"DOCX relationship from '{part}' references missing part '{target}'.")
            roles.setdefault(target, set()).add(relationship["role"])
            if target not in reachable:
                queue.append(target)
    return reachable, roles, external


def bundled_media_hashes() -> set[str]:
    with zipfile.ZipFile(DEFAULT_TEMPLATE, "r") as archive:
        return {
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.casefold().startswith("word/media/") and not name.endswith("/")
        }


def bundled_structural_attribute_allowlist() -> set[tuple[str, str, str, str]]:
    values: set[tuple[str, str, str, str]] = set()
    with zipfile.ZipFile(DEFAULT_TEMPLATE, "r") as archive:
        for part in archive.namelist():
            if not part.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(part))
            for node in root.iter():
                values.update((part, node.tag, attribute, value) for attribute, value in node.attrib.items())
    return values


def validate_package_no_secrets(archive: zipfile.ZipFile) -> None:
    _, _, external = package_relationship_graph(archive)
    parts = {
        part
        for part in archive.namelist()
        if part.endswith(".xml") or part.endswith(".rels")
    }
    for part in sorted(parts):
        if part not in archive.namelist() or not (part.endswith(".xml") or part.endswith(".rels")):
            continue
        root = ET.fromstring(archive.read(part))
        validate_secret_xml(root, part)
    for relationship in external:
        validate_no_secrets(
            relationship["target"],
            f"{relationship['relationship_part']}:Relationship.Target",
        )


def hidden_style_ids(archive: zipfile.ZipFile, style_parts: Iterable[str]) -> set[str]:
    hidden: set[str] = set()
    based_on: dict[str, str] = {}
    for part in style_parts:
        root = ET.fromstring(archive.read(part))
        for style in (node for node in root.iter() if node.tag.rsplit("}", 1)[-1].casefold() == "style"):
            style_id = style.get(f"{{{W_NS}}}styleId", "")
            if not style_id:
                continue
            if any(
                node.tag.rsplit("}", 1)[-1].casefold() in {"vanish", "webhidden", "specvanish"}
                for node in style.iter()
            ):
                hidden.add(style_id)
            base = next(
                (
                    node.get(f"{{{W_NS}}}val", "")
                    for node in style
                    if node.tag.rsplit("}", 1)[-1].casefold() == "basedon"
                ),
                "",
            )
            if base:
                based_on[style_id] = base
    changed = True
    while changed:
        changed = False
        for style_id, base in based_on.items():
            if base in hidden and style_id not in hidden:
                hidden.add(style_id)
                changed = True
    return hidden


def validate_customer_package(archive: zipfile.ZipFile) -> None:
    defaults, overrides = package_content_types(archive)
    reachable, roles, external = package_relationship_graph(archive)
    allowed_package_parts = reachable | {"[Content_Types].xml", "_rels/.rels"}
    allowed_package_parts.update(
        relationship_part_name(part)
        for part in reachable
        if relationship_part_name(part) in archive.namelist()
    )
    orphan_parts = [
        part
        for part in archive.namelist()
        if not part.endswith("/") and part not in allowed_package_parts
    ]
    if orphan_parts:
        raise ValueError(f"Customer DOCX contains unreferenced or unknown package part(s): {orphan_parts}.")
    all_relationships = [
        relationship
        for source in [None, *sorted(reachable)]
        for relationship in read_relationships(archive, source)
    ]
    for relationship in all_relationships:
        relationship_type = relationship["type"].casefold()
        if relationship["external"]:
            if relationship_type != HYPERLINK_REL_TYPE:
                raise ValueError(
                    f"Customer DOCX contains a forbidden external relationship type in "
                    f"{relationship['relationship_part']}."
                )
        elif relationship_type not in ALLOWED_INTERNAL_REL_TYPES:
            raise ValueError(
                f"Customer DOCX contains an unapproved relationship type in "
                f"{relationship['relationship_part']}."
            )
    forbidden_roles = {
        "afchunk",
        "attachedtemplate",
        "comments",
        "commentsextended",
        "commentsids",
        "control",
        "glossarydocument",
        "oleobject",
        "package",
        "people",
        "vbaproject",
    }
    forbidden_name_prefixes = (
        "customui/",
        "docprops/thumbnail",
        "word/activex/",
        "word/comments",
        "word/embeddings/",
        "word/glossary/",
        "word/people",
        "word/vbaproject",
    )
    for part in archive.namelist():
        lowered = part.casefold()
        content_type = part_content_type(part, defaults, overrides).casefold()
        if lowered.startswith(forbidden_name_prefixes) or any(
            marker in content_type
            for marker in ("activex", "comments", "macroenabled", "oleobject", "vbaproject", "vba-data")
        ):
            raise ValueError(f"Customer DOCX contains forbidden executable, comment, thumbnail, or hidden part '{part}'.")
    for part in reachable:
        if roles.get(part, set()) & forbidden_roles:
            raise ValueError(f"Customer DOCX contains forbidden relationship role(s) for '{part}': {roles[part]}.")
        if not part_content_type(part, defaults, overrides):
            raise ValueError(f"Customer DOCX reachable part '{part}' has no declared content type.")
    for relationship in external:
        role = relationship["role"]
        field = f"{relationship['relationship_part']}:Relationship.Target"
        validate_customer_text(relationship["target"], field)
        if role != "hyperlink":
            raise ValueError(f"Customer DOCX contains forbidden external '{role}' relationship in {field}.")
        target = relationship["target"]
        for _ in range(3):
            decoded_target = unquote(target)
            if decoded_target == target:
                break
            target = decoded_target
        parsed_target = urlsplit(target)
        if parsed_target.scheme.casefold() not in {"http", "https"} or not parsed_target.hostname:
            raise ValueError(f"Customer DOCX external hyperlink uses an unapproved URI scheme in {field}.")

    allowed_media = bundled_media_hashes()
    for part in archive.namelist():
        content_type = part_content_type(part, defaults, overrides).casefold()
        is_media = part.casefold().startswith("word/media/") or content_type.startswith(("image/", "audio/", "video/"))
        if is_media and not part.endswith("/"):
            digest = hashlib.sha256(archive.read(part)).hexdigest()
            if digest not in allowed_media:
                raise ValueError(f"Customer DOCX media part '{part}' is not in the bundled-template hash allowlist.")

    visible_roles = {"officedocument", "header", "footer", "footnotes", "endnotes"}
    visible_parts = {part for part in reachable if roles.get(part, set()) & visible_roles}
    style_parts = {part for part in reachable if "styles" in roles.get(part, set())}
    for part in style_parts:
        root = ET.fromstring(archive.read(part))
        if any(
            node.tag.rsplit("}", 1)[-1].casefold() in {"vanish", "webhidden", "specvanish"}
            for node in root.iter()
        ):
            raise ValueError(f"Customer DOCX styles contain hidden-text style/default markup in {part}.")
    hidden_styles = hidden_style_ids(archive, style_parts)
    if hidden_styles:
        raise ValueError(f"Customer DOCX styles contain hidden-text style(s): {sorted(hidden_styles)}.")
    forbidden_markup = {
        "altchunk",
        "commentrangeend",
        "commentrangestart",
        "commentreference",
        "del",
        "deltext",
        "ins",
        "movefrom",
        "movefromrangeend",
        "movefromrangestart",
        "moveto",
        "movetorangeend",
        "movetorangestart",
        "specvanish",
        "vanish",
        "webhidden",
    }
    for part in sorted(visible_parts):
        root = ET.fromstring(archive.read(part))
        if any(node.tag.rsplit("}", 1)[-1].casefold() in forbidden_markup for node in root.iter()):
            raise ValueError(f"Customer DOCX contains revisions, comments, hidden text, or external chunks in {part}.")
        for node in root.iter():
            if node.tag.rsplit("}", 1)[-1].casefold() in {"pstyle", "rstyle"}:
                if node.get(f"{{{W_NS}}}val", "") in hidden_styles:
                    raise ValueError(f"Customer DOCX uses hidden style '{node.get(f'{{{W_NS}}}val')}' in {part}.")
        validate_customer_xml_text(root, part)
        alt_text = [
            value
            for node in root.iter()
            for attribute, value in node.attrib.items()
            if attribute.rsplit("}", 1)[-1].casefold() in {"descr", "title"}
        ]
        for index, value in enumerate(alt_text):
            validate_customer_text(value, f"{part}:alt[{index}]")
    core = ET.fromstring(archive.read("docProps/core.xml"))
    for index, value in enumerate(node.text or "" for node in core.iter() if node.text):
        validate_customer_text(value, f"docProps/core.xml:text[{index}]")
    structural_allowlist = bundled_structural_attribute_allowlist()
    for part in sorted(allowed_package_parts):
        if not (part.endswith(".xml") or part.endswith(".rels")):
            continue
        root = ET.fromstring(archive.read(part))
        validate_customer_xml_text(root, part)
        for node in root.iter():
            for attribute, value in node.attrib.items():
                if (part, node.tag, attribute, value) not in structural_allowlist:
                    validate_customer_text(value, f"{part}:@{attribute.rsplit('}', 1)[-1]}")
    for source in [None, *sorted(reachable)]:
        for relationship in read_relationships(archive, source):
            validate_customer_text(
                relationship["target"],
                f"{relationship['relationship_part']}:Relationship.Target",
            )


def validate_package(path: Path, audience: str) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        corrupt = archive.testzip()
        if corrupt:
            raise ValueError(f"Generated DOCX contains a corrupt ZIP member: {corrupt}.")
        names = {name.replace("\\", "/").lower() for name in archive.namelist()}
        if (
            names & {"docprops/custom.xml", "docprops/app.xml"}
            or any(name.startswith("customxml/") for name in names)
            or any(name.startswith("docprops/thumbnail.") for name in names)
        ):
            raise ValueError("Generated DOCX still contains identity-bearing property parts.")
        settings = ET.fromstring(archive.read("word/settings.xml"))
        if any(
            element.tag.rsplit("}", 1)[-1].casefold() in {"docvars", "rsids", "docid"}
            for element in settings.iter()
        ):
            raise ValueError("Generated DOCX still contains document variables or document-tracking identifiers.")
        settings_bytes = archive.read("word/settings.xml").lower()
        if any(marker in settings_bytes for marker in (b"commondata", b"userid", b"hdid", b"wpscustom")):
            raise ValueError("Generated DOCX settings still contain WPS identity metadata.")
        for name in (item for item in archive.namelist() if item.endswith(".xml")):
            root = ET.fromstring(archive.read(name))
            for element in root.iter():
                if any(
                    attribute.rsplit("}", 1)[-1].casefold().startswith("rsid")
                    or attribute.rsplit("}", 1)[-1].casefold() in {"durableid", "paraid", "textid"}
                    for attribute in element.attrib
                ):
                    raise ValueError(f"Generated DOCX XML still contains paragraph/session tracking metadata in {name}.")
        validate_package_no_secrets(archive)
        if audience == "customer":
            validate_customer_package(archive)
    Document(str(path))


def resolved_path(path: str | Path) -> Path:
    return Path(path).expanduser().resolve()


def same_path(left: Path, right: Path) -> bool:
    return os.path.normcase(str(left)) == os.path.normcase(str(right))


def validate_paths(input_path: Path, template: Path, output: Path) -> None:
    if same_path(output, input_path):
        raise ValueError("Output path must not be the input JSON path.")
    if same_path(output, template):
        raise ValueError("Output path must not overwrite the template.")
    if output.suffix.lower() != ".docx":
        raise ValueError("Output path must use the .docx extension.")


def paragraph_number_id(paragraph: Paragraph) -> int | None:
    properties = paragraph._p.pPr
    if properties is None:
        return None
    number_properties = properties.find(qn("w:numPr"))
    if number_properties is None:
        return None
    number = number_properties.find(qn("w:numId"))
    if number is None:
        return None
    value = number.get(qn("w:val"))
    return int(value) if value is not None else None


def generated_content_targets(doc: Any) -> list[tuple[str, _Cell]]:
    info_table = unique_table_by_first_cell(doc, "故障基本信息")
    _, _, description = description_target(info_table)
    targets = [("fault_description", description)]
    for key, title in TEXT_BLOCKS.items():
        table = unique_table_by_first_cell(doc, title)
        targets.append((key, table.rows[1].cells[0]))
    return targets


def generated_field_paragraphs(doc: Any, data: dict[str, Any]) -> list[Paragraph]:
    paragraphs: list[Paragraph] = []
    for key in COVER_FIELDS:
        expected = plain_text(data[key])
        paragraphs.extend(paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == expected)
    info_table = unique_table_by_first_cell(doc, "故障基本信息")
    for _, label in INFO_FIELDS:
        _, _, target = label_target(info_table, label)
        paragraphs.extend(paragraph for paragraph in target.paragraphs if paragraph.text.strip())
    for _, cell in generated_content_targets(doc):
        paragraphs.extend(paragraph for paragraph in cell.paragraphs if paragraph.text.strip())
    return paragraphs


def text_runs(paragraph: Paragraph) -> list[Any]:
    return [run for run in paragraph.runs if run.text]


def line_spacing_is_15(paragraph: Paragraph) -> bool:
    spacing = paragraph.paragraph_format.line_spacing
    try:
        return abs(float(spacing) - 1.5) < 0.001
    except (TypeError, ValueError):
        return False


def audit_format_structure(path: Path, data: dict[str, Any]) -> dict[str, int]:
    doc = Document(str(path))
    failures: list[str] = []
    metrics = {
        "ordinary_all_bold_paragraphs": 0,
        "mixed_bold_paragraphs": 0,
        "native_numbered_paragraphs": 0,
        "duplicate_number_paragraphs": 0,
    }

    expected_line_bold = {
        plain_text(data["report_title"]),
        plain_text(data["fault_level"]),
        plain_text(data["fault_start_time"]),
        plain_text(data["business_recovery_time"]),
    }
    for field, _ in generated_content_targets(doc):
        for line in format_items(data[field], field):
            if line.line_bold:
                expected_line_bold.add(plain_text(line.text))

    for field, cell in generated_content_targets(doc):
        expected = format_items(data[field], field)
        actual = [paragraph for paragraph in cell.paragraphs if paragraph.text.strip()]
        if len(actual) != len(expected):
            failures.append(
                f"field '{field}' rendered {len(actual)} paragraph(s), expected {len(expected)}"
            )
            continue
        for index, (paragraph, line) in enumerate(zip(actual, expected), 1):
            number_id = paragraph_number_id(paragraph)
            if line.numbered and number_id is None:
                failures.append(f"field '{field}' paragraph {index} lost native numbering")
            if not line.numbered and number_id is not None:
                failures.append(f"field '{field}' paragraph {index} has unexpected native numbering")
            if number_id is not None:
                metrics["native_numbered_paragraphs"] += 1

    for paragraph in generated_field_paragraphs(doc, data):
        text = paragraph.text.strip()
        if not line_spacing_is_15(paragraph):
            failures.append(f"paragraph does not use 1.5 line spacing: {text[:80]}")
        for name, spacing in (
            ("space_before", paragraph.paragraph_format.space_before),
            ("space_after", paragraph.paragraph_format.space_after),
        ):
            if spacing is not None and spacing.pt != 0:
                failures.append(f"paragraph {name} must be 0 pt: {text[:80]}")
        if DUPLICATE_NUMBER_RE.match(plain_text(text)):
            metrics["duplicate_number_paragraphs"] += 1
            failures.append(f"paragraph contains duplicate numbering: {text[:80]}")
        runs = text_runs(paragraph)
        if not runs:
            continue
        bold_values = [run.bold is True for run in runs]
        all_bold = all(bold_values)
        if any(bold_values) and not all_bold:
            metrics["mixed_bold_paragraphs"] += 1
        if all_bold and plain_text(text) not in expected_line_bold and not LABEL_RE.fullmatch(plain_text(text)):
            metrics["ordinary_all_bold_paragraphs"] += 1
            failures.append(f"ordinary paragraph is entirely bold: {text[:80]}")
        if plain_text(text) in expected_line_bold and not all_bold:
            failures.append(f"explicit line-bold paragraph is not entirely bold: {text[:80]}")

    document_text = "\n".join(paragraph.text for paragraph in iter_all_paragraphs(doc))
    for placeholder in TEMPLATE_PLACEHOLDERS:
        if placeholder in document_text:
            failures.append(f"template placeholder remains: {placeholder}")
    if failures:
        raise ValueError("DOCX format audit failed: " + "; ".join(failures))
    return metrics


def atomic_save(doc: Any, output: Path, audience: str, data: dict[str, Any]) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    rendered: Path | None = None
    sanitized: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent, prefix=f".{output.name}.", suffix=".rendered", delete=False) as handle:
            rendered = Path(handle.name)
        with tempfile.NamedTemporaryFile(dir=output.parent, prefix=f".{output.name}.", suffix=".sanitized", delete=False) as handle:
            sanitized = Path(handle.name)
        doc.save(str(rendered))
        sanitize_package(rendered, sanitized)
        validate_package(sanitized, audience)
        audit_format_structure(sanitized, data)
        os.replace(sanitized, output)
        sanitized = None
    finally:
        for temporary in (rendered, sanitized):
            if temporary is not None:
                temporary.unlink(missing_ok=True)


def generate(data: dict[str, Any], template: Path, output: Path) -> None:
    validated = validate_input(data)
    if not template.is_file():
        raise FileNotFoundError(f"Template not found: {template}")
    doc = Document(str(template))
    update_document(doc, validated)
    apply_line_spacing(doc)
    mark_fields_for_update(doc)
    clean_core_properties(doc, validated["report_title"], "ZStack fault analysis report")
    atomic_save(doc, output, validated["audience"], validated)


def reject_duplicate_object(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Input JSON contains duplicate key '{key}'.")
        result[key] = value
    return result


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as handle:
        return validate_input(json.load(handle, object_pairs_hook=reject_duplicate_object))


def print_template(template: Path) -> None:
    if not template.is_file():
        raise FileNotFoundError(f"Template not found: {template}")
    doc = Document(str(template))
    summary = probe_template(doc)
    summary["template"] = str(template)
    json.dump(summary, sys.stdout, ensure_ascii=True, indent=2)
    sys.stdout.write("\n")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_json", nargs="?", help="UTF-8 or UTF-8-BOM JSON input file.")
    parser.add_argument("--out", help="Output DOCX path.")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Template DOCX path.")
    parser.add_argument("--print-template", action="store_true", help="Probe and print writable template fields as JSON.")
    args = parser.parse_args(argv)

    try:
        template = resolved_path(args.template)
        if args.print_template:
            if args.input_json or args.out:
                raise ValueError("--print-template cannot be combined with input_json or --out.")
            print_template(template)
            return 0
        if not args.input_json or not args.out:
            parser.error("input_json and --out are required unless --print-template is used.")
        input_path = resolved_path(args.input_json)
        output = resolved_path(args.out)
        validate_paths(input_path, template, output)
        generate(load_json(input_path), template, output)
        print(f"Generated: {output}")
        return 0
    except (OSError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
