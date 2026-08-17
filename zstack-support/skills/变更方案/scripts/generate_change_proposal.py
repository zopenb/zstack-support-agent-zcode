#!/usr/bin/env python3
"""Render validated ZStack change-proposal JSON into the standard DOCX template."""

from __future__ import annotations

import sys

if sys.version_info < (3, 10):
    print("ERROR: this generator requires Python 3.10 or newer.", file=sys.stderr)
    raise SystemExit(2)

import argparse
import base64
import binascii
from copy import deepcopy
import hashlib
import ipaddress
import json
import os
from pathlib import Path
import posixpath
import re
import tempfile
from typing import Any, Iterable
import zipfile
from xml.etree import ElementTree as ET
from urllib.parse import unquote, urlsplit


SCRIPT_PATH = Path(__file__).resolve()
SKILL_DIR = SCRIPT_PATH.parents[1]
PLUGIN_DIR = SCRIPT_PATH.parents[3]
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "2021XX-XX项目XX问题-ZStack变更方案模板v1.2.docx"
CHECKLIST_TITLE = "ZStack云计算运维风险checklist"
CHECKLIST_MODULE_CONTRACT = (
    ("cloud-platform-unmount", "云平台\n卸载", "高"),
    ("cloud-platform-delete", "云平台\n删除", "高"),
    ("host-layer-change", "底层\n变更", "高"),
    ("distributed-storage-high", "分布式\n存储\n高风险", "高"),
    ("distributed-storage-medium", "分布式\n存储\n中风险", "中"),
    ("network", "网络", "高"),
    ("other-medium-high", "其他\n中高\n风险", "高"),
)


def dependency_directories() -> list[Path]:
    override = os.environ.get("ZSTACK_SUPPORT_PYTHONPATH")
    if override:
        return [Path(item).expanduser().resolve() for item in override.split(os.pathsep) if item]
    if sys.platform == "win32":
        base = Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local"))
        return [base / "ZStackSupportAgent" / "python"]
    if sys.platform == "darwin":
        return [Path.home() / "Library" / "Caches" / "ZStackSupportAgent" / "python"]
    return [Path.home() / ".cache" / "ZStackSupportAgent" / "python"]


DEPENDENCY_DIRS = dependency_directories()
for dependency_dir in reversed(DEPENDENCY_DIRS):
    if str(dependency_dir) not in sys.path:
        sys.path.insert(0, str(dependency_dir))

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.table import _Cell
    from docx.text.paragraph import Paragraph
    from lxml import etree as LET
except (ImportError, ModuleNotFoundError) as exc:
    target = DEPENDENCY_DIRS[0]
    requirements = PLUGIN_DIR / "requirements.txt"
    print(
        "ERROR: python-docx is unavailable. Install the plugin's pinned dependencies with:\n"
        f'  "{sys.executable}" -m pip install --target "{target}" -r "{requirements}"\n'
        "Set ZSTACK_SUPPORT_PYTHONPATH to use a different private dependency directory.",
        file=sys.stderr,
    )
    raise SystemExit(2) from exc


COVER_FIELDS = {
    "cover_title": "ZStack运维变更方案",
    "change_title": "XX变更",
    "risk_level": "风险等级：X",
    "document_date": "时间：XX.XX.XX",
}

CONFIG_FIELDS = {
    "software_info": "软件信息",
    "hardware_config": "硬件配置",
    "business_info": "业务信息",
}

SECTION_MAP = {
    "overview": ("变更概述", "变更步骤"),
    "change_principles": ("2.2 变更原则及变更范围", "2.3 变更整体流程"),
    "overall_flow": ("2.3 变更整体流程", "2.4 变更具体步骤"),
    "detailed_steps": ("2.4 变更具体步骤", "风险评估"),
    "risk_mitigations": ("3.2风险预案", "回退方案"),
    "rollback_plan": ("回退方案", "紧急预案"),
    "emergency_plan": ("紧急预案", "变更计划"),
}

SCALAR_REQUIRED = {"audience", *COVER_FIELDS, *CONFIG_FIELDS}
CONTENT_REQUIRED = {*SECTION_MAP, "risks"}
PLAN_FIELDS = {"change_time", "maintenance_window", "executor", "executor_phone", "supervisor", "supervisor_phone"}
CHECKLIST_FIELDS = {"checklist_items", "checklist_decisions", "mark_unmatched_checklist_as_no"}
ALLOWED_FIELDS = (
    SCALAR_REQUIRED
    | CONTENT_REQUIRED
    | PLAN_FIELDS
    | CHECKLIST_FIELDS
    | {"change_plan", "customer_identity_handling"}
)
AUDIENCES = {"internal", "customer"}
CUSTOMER_IDENTITY_HANDLING = {"confirmed_same_customer", "generalized"}
IMPORTANT_LABELS = {
    "执行动作",
    "变更动作",
    "回退动作",
    "验证动作",
    "风险等级",
    "变更时间",
    "变更窗口",
    "变更执行人",
    "变更监督人",
}

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CORE_NS = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
DC_NS = "http://purl.org/dc/elements/1.1/"
DCTERMS_NS = "http://purl.org/dc/terms/"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL_BASE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
HYPERLINK_REL_TYPE = f"{OFFICE_REL_BASE}/hyperlink".casefold()
ALLOWED_INTERNAL_REL_TYPES = {
    "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties",
    *(f"{OFFICE_REL_BASE}/{role}" for role in (
        "officeDocument", "fontTable", "footer", "header", "image", "numbering", "settings",
        "styles", "theme", "endnotes", "footnotes", "webSettings",
    )),
}
ALLOWED_INTERNAL_REL_TYPES = {value.casefold() for value in ALLOWED_INTERNAL_REL_TYPES}
URL_RE = re.compile(r"https?://[^\s<>()\[\]{}\"']+", re.IGNORECASE)
DOTTED_IPV4_RE = re.compile(
    r"(?<![A-Za-z0-9_.])(?:0x[0-9a-f]+|\d+)(?:\.(?:0x[0-9a-f]+|\d+)){1,3}(?![A-Za-z0-9_.])",
    re.IGNORECASE,
)
INTEGER_IPV4_RE = re.compile(r"(?<![A-Za-z0-9_.])(?:0x[0-9a-f]{1,8}|0[0-7]{7,11}|\d{8,10})(?![A-Za-z0-9_.])", re.IGNORECASE)
IPV6_RE = re.compile(r"(?<![0-9a-f:])(?:[0-9a-f]{0,4}:){2,7}[0-9a-f]{0,4}(?![0-9a-f:])", re.IGNORECASE)
UUID_RE = re.compile(r"(?<![0-9a-f])[0-9a-f]{8}(?:-[0-9a-f]{4}){3}-[0-9a-f]{12}(?![0-9a-f])", re.IGNORECASE)
COMPACT_GUID_RE = re.compile(r"(?<![0-9a-f])[0-9a-f]{12}[1-8][0-9a-f]{3}[89ab][0-9a-f]{15}(?![0-9a-f])", re.IGNORECASE)
INTERNAL_TICKET_RE = re.compile(r"(?<![A-Za-z0-9_])(?:TIC|BUG|SUG|JIRA)-\d+(?![A-Za-z0-9_])", re.IGNORECASE)
RELATIVE_INTERNAL_LINK_RE = re.compile(
    r"(?:forum\.php\?[^\s]*\btid=|/(?:browse|issues?)/|(?:^|[/])(?:bbs|jira|confluence)(?:[/])|confluence/(?:pages|display)/)",
    re.IGNORECASE,
)
INTERNAL_HOST_RE = re.compile(
    r"(?<![A-Za-z0-9_.-])(?:localhost|[A-Za-z0-9_-]+(?:\.[A-Za-z0-9_-]+)*\.(?:internal|local|corp)|(?:bbs|jira|confluence)(?:[.-][A-Za-z0-9_-]+)+)(?::\d+)?",
    re.IGNORECASE,
)
PEM_PRIVATE_KEY_RE = re.compile(
    r"-----BEGIN (?:(?:RSA|EC|DSA|OPENSSH|ENCRYPTED) )?PRIVATE KEY-----|"
    r"-----BEGIN PGP PRIVATE KEY BLOCK-----",
    re.IGNORECASE,
)
AUTHORIZATION_SECRET_RE = re.compile(
    r"(?<![A-Za-z0-9_])Authorization\s*[:=：]\s*(?:Bearer|Basic)\s+[^\s,;]+",
    re.IGNORECASE,
)
BASIC_CREDENTIAL_RE = re.compile(r"(?<![A-Za-z0-9_])Basic\s+([A-Za-z0-9+/]+={0,2})(?![A-Za-z0-9+/=])", re.IGNORECASE)
SECRET_ASSIGNMENT_RE = re.compile(
    r"(?<![A-Za-z0-9_])(?:password|passwd|pwd|client[_-]?secret|secret|api[_-]?key|access[_-]?key|auth[_-]?token|token)(?![A-Za-z0-9_])\s*[:=：]\s*[\"']?([^\s\"',;；]+)",
    re.IGNORECASE,
)
URI_USERINFO_RE = re.compile(r"https?://[^/\s:@]+:[^@\s/]+@", re.IGNORECASE)
KNOWN_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9_])(?:AKIA[0-9A-Z]{16}|ghp_[A-Za-z0-9]{20,}|github_pat_[A-Za-z0-9_]{20,})(?![A-Za-z0-9_])")
JWT_RE = re.compile(r"(?<![A-Za-z0-9_])eyJ[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{8,}(?![A-Za-z0-9_])")


def plain_text(value: str) -> str:
    return value.replace("**", "").strip()


def normalize(value: str) -> str:
    return re.sub(r"\s+", "", value or "").casefold()


def validate_nonempty_string(value: Any, field: str) -> str:
    if not isinstance(value, str):
        raise ValueError(f"Field '{field}' must be a string.")
    value = value.strip()
    if not value or not plain_text(value):
        raise ValueError(f"Field '{field}' must not be empty; use an explicit value such as '待补充'.")
    return value


def validate_content(value: Any, field: str) -> str | list[str]:
    if isinstance(value, str):
        validated = validate_nonempty_string(value, field)
        if not [line for line in validated.splitlines() if line.strip()]:
            raise ValueError(f"Field '{field}' must contain text.")
        return validated
    if not isinstance(value, list):
        raise ValueError(f"Field '{field}' must be a string or a non-empty array of strings.")
    if not value:
        raise ValueError(f"Field '{field}' must not be an empty array.")
    return [validate_nonempty_string(item, f"{field}[{index}]") for index, item in enumerate(value)]


def iter_text_values(value: Any, path: str = "$") -> Iterable[tuple[str, str]]:
    if isinstance(value, str):
        yield path, value
    elif isinstance(value, dict):
        for key, item in value.items():
            yield from iter_text_values(item, f"{path}.{key}")
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from iter_text_values(item, f"{path}[{index}]")


def decoded_text(text: str) -> str:
    values = [text]
    current = text
    for _ in range(3):
        decoded = unquote(current)
        if decoded == current:
            break
        values.append(decoded)
        current = decoded
    return "\n".join(values)


def validate_no_secrets(text: str, field: str) -> None:
    scan = decoded_text(text)
    if PEM_PRIVATE_KEY_RE.search(scan):
        raise ValueError(f"Output field '{field}' contains private-key PEM material.")
    if AUTHORIZATION_SECRET_RE.search(scan) or URI_USERINFO_RE.search(scan):
        raise ValueError(f"Output field '{field}' contains an authorization credential.")
    for match in BASIC_CREDENTIAL_RE.finditer(scan):
        try:
            decoded = base64.b64decode(match.group(1), validate=True)
        except (binascii.Error, ValueError):
            continue
        if b":" in decoded:
            raise ValueError(f"Output field '{field}' contains an authorization credential.")
    if KNOWN_TOKEN_RE.search(scan) or JWT_RE.search(scan):
        raise ValueError(f"Output field '{field}' contains an access token.")
    safe_values = {
        "待补充", "待确认", "已脱敏", "redacted", "<redacted>", "***", "none", "null",
        "disabled", "unset", "unavailable", "omitted", "false", "off", "not-set", "not_configured",
    }
    for match in SECRET_ASSIGNMENT_RE.finditer(scan):
        if match.group(1).strip().casefold() not in {value.casefold() for value in safe_values}:
            raise ValueError(f"Output field '{field}' contains an assigned password, token, or secret.")


def parse_ipv4_number(value: str) -> int | None:
    try:
        if value.casefold().startswith("0x"):
            number = int(value, 16)
        elif len(value) > 1 and value.startswith("0"):
            number = int(value, 8) if all(character in "01234567" for character in value) else int(value, 10)
        else:
            number = int(value, 10)
    except ValueError:
        return None
    return number if 0 <= number <= 0xFFFFFFFF else None


def parse_legacy_ipv4(value: str) -> int | None:
    components = value.split(".")
    if not 1 <= len(components) <= 4 or any(not component for component in components):
        return None
    numbers = [parse_ipv4_number(component) for component in components]
    if any(number is None for number in numbers):
        return None
    parsed = [int(number) for number in numbers if number is not None]
    limits = {
        1: (0xFFFFFFFF,),
        2: (0xFF, 0xFFFFFF),
        3: (0xFF, 0xFF, 0xFFFF),
        4: (0xFF, 0xFF, 0xFF, 0xFF),
    }[len(parsed)]
    if any(number > limit for number, limit in zip(parsed, limits)):
        return None
    result = parsed[-1]
    shift = 8 * (4 - len(parsed) + 1)
    for number in reversed(parsed[:-1]):
        result |= number << shift
        shift += 8
    return result


def version_context(text: str, start: int) -> bool:
    prefix = text[max(0, start - 32):start]
    return bool(re.search(r"(?i)(?:version|ver|release|build|版本|构建)\s*[:=v-]?\s*$", prefix))


def semantic_version_candidate(value: str) -> bool:
    parts = value.split(".")
    return (
        len(parts) == 3
        and all(part.isdigit() and (part == "0" or not part.startswith("0")) for part in parts)
        and int(parts[0]) <= 20
        and all(int(part) <= 99 for part in parts[1:])
    )


def ip_number_context(text: str, start: int, end: int) -> bool:
    prefix = text[max(0, start - 32):start]
    suffix = text[end:min(len(text), end + 12)]
    return bool(
        re.search(
            r"(?i)(?:(?<![A-Za-z0-9_])(?:ip(?:v4)?|host|endpoint)(?![A-Za-z0-9_])|地址|端点)"
            r"\s*[:=：]?\s*$",
            prefix,
        )
        or re.match(r"^:\d{1,5}\b", suffix)
    )


def validate_customer_text(text: str, field: str) -> None:
    text = decoded_text(text)
    lowered = text.casefold()
    if UUID_RE.search(text) or COMPACT_GUID_RE.search(text):
        raise ValueError(f"Customer output field '{field}' contains a raw GUID; generalize the identifier.")
    if INTERNAL_TICKET_RE.search(text):
        raise ValueError(f"Customer output field '{field}' contains an internal ticket identifier.")
    if RELATIVE_INTERNAL_LINK_RE.search(URL_RE.sub("", text)):
        raise ValueError(f"Customer output field '{field}' contains a relative internal-system link.")
    if "file://" in lowered or re.search(r"\\\\[^\\\s]+\\[^\s]+", text) or re.search(r"\b[A-Za-z]:\\", text):
        raise ValueError(f"Customer output field '{field}' contains an internal file endpoint.")
    if INTERNAL_HOST_RE.search(text):
        raise ValueError(f"Customer output field '{field}' contains an internal BBS/Jira/Confluence or private host.")
    for match in DOTTED_IPV4_RE.finditer(text):
        candidate = match.group(0)
        if (
            parse_legacy_ipv4(candidate) is None
            or version_context(text, match.start())
            or semantic_version_candidate(candidate)
            or field.casefold().endswith(".version")
        ):
            continue
        if candidate.count(".") == 1 and not (
            candidate.casefold().startswith("0x")
            or candidate.startswith("0")
            or ip_number_context(text, match.start(), match.end())
        ):
            continue
        raise ValueError(f"Customer output field '{field}' contains a raw IP address; generalize the endpoint.")
    for match in INTEGER_IPV4_RE.finditer(text):
        candidate = match.group(0)
        encoded_ip_form = (
            candidate.casefold().startswith("0x") and len(candidate[2:]) == 8
        ) or (
            candidate.startswith("0") and len(candidate) >= 8 and all(character in "01234567" for character in candidate)
        )
        if parse_ipv4_number(candidate) is not None and (
            encoded_ip_form or ip_number_context(text, match.start(), match.end())
        ):
            raise ValueError(f"Customer output field '{field}' contains an integer-form IPv4 address.")
    for match in IPV6_RE.finditer(text):
        candidate = match.group(0)
        try:
            ipaddress.ip_address(candidate)
        except ValueError:
            continue
        raise ValueError(f"Customer output field '{field}' contains a raw IPv6 address; generalize the endpoint.")
    for match in URL_RE.finditer(text):
        url = match.group(0).rstrip(".,;:!?)]}，。；：！？")
        parsed = urlsplit(url)
        host = (parsed.hostname or "").casefold()
        target = f"{host}{parsed.path}".casefold()
        if any(token in target for token in ("bbs", "jira", "confluence", "atlassian.net")):
            raise ValueError(f"Customer output field '{field}' contains an internal support-system link.")
        try:
            address = ipaddress.ip_address(host)
        except ValueError:
            address = None
        if address is not None or parse_legacy_ipv4(host) is not None:
            raise ValueError(f"Customer output field '{field}' contains a raw IP URL endpoint.")


def validate_customer_content(data: dict[str, Any]) -> None:
    for field, text in iter_text_values(data):
        validate_customer_text(text, field)


def validate_secret_content(data: dict[str, Any]) -> None:
    for field, text in iter_text_values(data):
        validate_no_secrets(text, field)


def validate_customer_xml_text(root: ET.Element, field: str) -> None:
    text_nodes = [
        node.text or ""
        for node in root.iter()
        if node.tag.rsplit("}", 1)[-1].casefold() in {"t", "instrtext"} and (node.text or "")
    ]
    for index, value in enumerate(text_nodes):
        validate_customer_text(value, f"{field}:text[{index}]")
    if text_nodes:
        validate_customer_text("\n".join(text_nodes), field)
    for index, paragraph in enumerate(
        node for node in root.iter() if node.tag.rsplit("}", 1)[-1].casefold() == "p"
    ):
        rendered = "".join(
            node.text or ""
            for node in paragraph.iter()
            if node.tag.rsplit("}", 1)[-1].casefold() in {"t", "instrtext"}
        )
        if rendered:
            validate_customer_text(rendered, f"{field}:paragraph[{index}]")


def validate_secret_xml(root: ET.Element, field: str) -> None:
    values = [node.text or "" for node in root.iter() if node.text]
    values.extend(value for node in root.iter() for value in node.attrib.values() if value)
    for index, value in enumerate(values):
        validate_no_secrets(value, f"{field}:value[{index}]")
    if values:
        validate_no_secrets("\n".join(values), field)
        validate_no_secrets("".join(values), field)
    for index, paragraph in enumerate(
        node for node in root.iter() if node.tag.rsplit("}", 1)[-1].casefold() == "p"
    ):
        rendered = "".join(node.text or "" for node in paragraph.iter() if node.text)
        if rendered:
            validate_no_secrets(rendered, f"{field}:paragraph[{index}]")


def validate_selector(value: Any, field: str, allow_string: bool) -> str | dict[str, str]:
    if isinstance(value, str):
        if not allow_string:
            raise ValueError(f"Field '{field}' must be an object with an 'involved' boolean.")
        return validate_nonempty_string(value, field)
    if not isinstance(value, dict):
        expected = "a string or selector object" if allow_string else "a selector object"
        raise ValueError(f"Field '{field}' must be {expected}.")
    allowed = {"row_id", "module_id", "module", "operation", "impact", "level"}
    unknown = sorted(set(value) - allowed)
    if unknown:
        raise ValueError(f"Unknown selector field(s) in '{field}': {', '.join(unknown)}.")
    if not value.get("row_id") and not value.get("operation"):
        raise ValueError(f"Field '{field}' must identify a checklist row by row_id or operation.")
    selector = {key: validate_nonempty_string(item, f"{field}.{key}") for key, item in value.items()}
    if "row_id" not in selector and ("module" in selector or "impact" in selector) and "operation" not in selector:
        raise ValueError(f"Field '{field}' requires operation when matching by module or impact.")
    return selector


def validate_checklist_items(value: Any) -> list[str | dict[str, str]]:
    if not isinstance(value, list):
        raise ValueError("Field 'checklist_items' must be an array.")
    return [validate_selector(item, f"checklist_items[{index}]", allow_string=True) for index, item in enumerate(value)]


def validate_checklist_decisions(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        raise ValueError("Field 'checklist_decisions' must be an array.")
    decisions: list[dict[str, Any]] = []
    for index, item in enumerate(value):
        field = f"checklist_decisions[{index}]"
        if not isinstance(item, dict):
            raise ValueError(f"Field '{field}' must be an object.")
        allowed = {"row_id", "module_id", "module", "operation", "impact", "level", "involved", "reason"}
        unknown = sorted(set(item) - allowed)
        if unknown:
            raise ValueError(f"Unknown decision field(s) in '{field}': {', '.join(unknown)}.")
        if "involved" not in item or type(item["involved"]) is not bool:
            raise ValueError(f"Field '{field}.involved' must be a JSON boolean.")
        selector_input = {
            key: item[key]
            for key in ("row_id", "module_id", "module", "operation", "impact", "level")
            if key in item
        }
        selector = validate_selector(selector_input, field, allow_string=False)
        assert isinstance(selector, dict)
        decision: dict[str, Any] = {**selector, "involved": item["involved"]}
        if "reason" in item:
            decision["reason"] = validate_nonempty_string(item["reason"], f"{field}.reason")
        decisions.append(decision)
    return decisions


def validate_input(data: Any) -> dict[str, Any]:
    if not isinstance(data, dict):
        raise ValueError("Input JSON must be an object.")
    unknown = sorted(set(data) - ALLOWED_FIELDS)
    missing = sorted((SCALAR_REQUIRED | CONTENT_REQUIRED) - set(data))
    if unknown:
        raise ValueError(f"Unknown input field(s): {', '.join(unknown)}.")
    if missing:
        raise ValueError(f"Missing required field(s): {', '.join(missing)}.")

    has_change_plan = "change_plan" in data
    structured_plan = PLAN_FIELDS & set(data)
    if has_change_plan and structured_plan:
        raise ValueError("Use either 'change_plan' or the structured change-plan fields, not both.")
    if not has_change_plan:
        plan_missing = sorted(PLAN_FIELDS - set(data))
        if plan_missing:
            raise ValueError(f"Missing structured change-plan field(s): {', '.join(plan_missing)}.")
    if "checklist_items" not in data and "checklist_decisions" not in data:
        raise ValueError("One of 'checklist_items' or 'checklist_decisions' is required, and may be an empty array.")

    validated: dict[str, Any] = {}
    audience = validate_nonempty_string(data["audience"], "audience")
    if audience not in AUDIENCES:
        raise ValueError("Field 'audience' must be 'internal' or 'customer'.")
    validated["audience"] = audience
    for field in sorted(SCALAR_REQUIRED - {"audience"}):
        validated[field] = validate_nonempty_string(data[field], field)
    for field in sorted(CONTENT_REQUIRED):
        validated[field] = validate_content(data[field], field)
    if has_change_plan:
        validated["change_plan"] = validate_content(data["change_plan"], "change_plan")
    else:
        for field in sorted(PLAN_FIELDS):
            validated[field] = validate_nonempty_string(data[field], field)
    if "checklist_items" in data:
        validated["checklist_items"] = validate_checklist_items(data["checklist_items"])
    if "checklist_decisions" in data:
        validated["checklist_decisions"] = validate_checklist_decisions(data["checklist_decisions"])
    if "mark_unmatched_checklist_as_no" in data:
        if type(data["mark_unmatched_checklist_as_no"]) is not bool:
            raise ValueError("Field 'mark_unmatched_checklist_as_no' must be a JSON boolean.")
        validated["mark_unmatched_checklist_as_no"] = data["mark_unmatched_checklist_as_no"]
    if audience == "customer":
        if "customer_identity_handling" not in data:
            raise ValueError(
                "Customer output requires 'customer_identity_handling' to be "
                "'confirmed_same_customer' or 'generalized'."
            )
        identity_handling = validate_nonempty_string(
            data["customer_identity_handling"], "customer_identity_handling"
        )
        if identity_handling not in CUSTOMER_IDENTITY_HANDLING:
            raise ValueError(
                "Field 'customer_identity_handling' must be 'confirmed_same_customer' or 'generalized'."
            )
        validated["customer_identity_handling"] = identity_handling
        validate_customer_content(validated)
    elif "customer_identity_handling" in data:
        raise ValueError("Field 'customer_identity_handling' is only valid when audience is 'customer'.")
    validate_secret_content(validated)
    return validated


def format_items(value: str | list[str]) -> list[str]:
    source = value if isinstance(value, list) else [line.strip() for line in value.splitlines() if line.strip()]
    lines: list[str] = []
    for index, item in enumerate(source, 1):
        if re.match(r"^\d+[.、]", item) or re.match(r"^[^:：]{1,24}[:：]", plain_text(item)):
            lines.append(item)
        else:
            lines.append(f"{index}. {item}")
    return lines


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def run_prototype(paragraph: Paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    ppr = paragraph._p.pPr
    paragraph_rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
    if paragraph_rpr is not None:
        return deepcopy(paragraph_rpr)
    return None


def add_run(paragraph: Paragraph, text: str, prototype: Any, bold: bool = False) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    if prototype is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, deepcopy(prototype))
    if bold:
        run.bold = True


def write_rich_text(paragraph: Paragraph, text: str, force_bold: bool = False, prototype: Any = None) -> None:
    prototype = deepcopy(prototype) if prototype is not None else run_prototype(paragraph)
    clear_paragraph(paragraph)
    paragraph.paragraph_format.line_spacing = 1.5
    segments = re.split(r"(\*\*.*?\*\*)", text)
    at_start = True
    whole_line_bold = force_bold
    for segment in segments:
        if not segment:
            continue
        marked = segment.startswith("**") and segment.endswith("**") and len(segment) >= 4
        if marked:
            segment = segment[2:-2]
        if not segment:
            continue
        if at_start and not marked:
            match = re.match(r"^([^:：]{1,24}[:：])(.*)$", segment, flags=re.DOTALL)
            if match:
                label = match.group(1)
                label_name = label[:-1].strip()
                whole_line_bold = force_bold or label_name in IMPORTANT_LABELS
                add_run(paragraph, label, prototype, bold=True)
                add_run(
                    paragraph,
                    match.group(2),
                    prototype,
                    bold=whole_line_bold,
                )
                at_start = False
                continue
        add_run(paragraph, segment, prototype, bold=whole_line_bold or marked)
        at_start = False


def paragraph_prototype(paragraph: Paragraph) -> tuple[Any, Any]:
    return (
        deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None,
        run_prototype(paragraph),
    )


def apply_paragraph_prototype(paragraph: Paragraph, prototype: tuple[Any, Any]) -> None:
    ppr, _ = prototype
    if paragraph._p.pPr is not None:
        paragraph._p.remove(paragraph._p.pPr)
    if ppr is not None:
        paragraph._p.insert(0, deepcopy(ppr))


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def set_cell_text(cell: _Cell, text: str, force_bold: bool = False) -> None:
    paragraphs = list(cell.paragraphs)
    prototype = paragraph_prototype(paragraphs[0])
    for paragraph in paragraphs[1:]:
        remove_paragraph(paragraph)
    apply_paragraph_prototype(paragraphs[0], prototype)
    write_rich_text(paragraphs[0], text, force_bold=force_bold, prototype=prototype[1])


def unique_table_by_first_cell(doc: Any, first_cell_text: str):
    matches = [
        table
        for table in doc.tables
        if table.rows and table.rows[0].cells and table.rows[0].cells[0].text.strip() == first_cell_text
    ]
    if len(matches) != 1:
        raise ValueError(f"Template must contain exactly one table titled '{first_cell_text}', found {len(matches)}.")
    return matches[0]


def unique_paragraph(doc: Any, text: str) -> tuple[int, Paragraph]:
    matches = [(index, paragraph) for index, paragraph in enumerate(doc.paragraphs) if paragraph.text.strip() == text]
    if len(matches) != 1:
        raise ValueError(f"Template must contain exactly one paragraph '{text}', found {len(matches)}.")
    return matches[0]


def body_paragraphs(doc: Any, heading_text: str, next_heading_text: str | None) -> list[Paragraph]:
    start_index, _ = unique_paragraph(doc, heading_text)
    end_index = len(doc.paragraphs)
    if next_heading_text is not None:
        end_index, _ = unique_paragraph(doc, next_heading_text)
    if start_index >= end_index:
        raise ValueError(f"Template section '{heading_text}' is not before '{next_heading_text}'.")
    body = list(doc.paragraphs[start_index + 1 : end_index])
    if not body:
        raise ValueError(f"Template section '{heading_text}' has no writable body paragraph.")
    return body


def replace_section(doc: Any, heading_text: str, next_heading_text: str | None, lines: list[str]) -> None:
    _, heading = unique_paragraph(doc, heading_text)
    body = body_paragraphs(doc, heading_text, next_heading_text)
    prototype = paragraph_prototype(body[0])
    for paragraph in body:
        remove_paragraph(paragraph)
    anchor = heading
    for line in lines:
        element = OxmlElement("w:p")
        anchor._p.addnext(element)
        paragraph = Paragraph(element, anchor._parent)
        apply_paragraph_prototype(paragraph, prototype)
        write_rich_text(paragraph, line, prototype=prototype[1])
        anchor = paragraph


def config_target(table: Any, label: str) -> tuple[int, int, _Cell]:
    matches: list[tuple[int, int, _Cell]] = []
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells[:-1]):
            if cell.text.strip() == label and row.cells[column_index + 1]._tc is not cell._tc:
                matches.append((row_index, column_index + 1, row.cells[column_index + 1]))
    if len(matches) != 1:
        raise ValueError(f"Template configuration label '{label}' has {len(matches)} writable targets.")
    return matches[0]


def stable_row_id(module_id: str, operation: str, impact: str, level: str) -> str:
    material = "|".join(normalize(item) for item in (module_id, operation, impact, level)).encode("utf-8")
    return f"cl-{hashlib.sha256(material).hexdigest()[:12]}"


def checklist_rows(doc: Any) -> list[dict[str, Any]]:
    table = unique_table_by_first_cell(doc, CHECKLIST_TITLE)
    if len(table.rows) < 3:
        raise ValueError("Template checklist has no data rows.")
    expected_header = ["模块", "操作", "影响风险", "等级", "是否涉及"]
    actual_header = [cell.text.strip() for cell in table.rows[1].cells[:5]]
    if actual_header != expected_header:
        raise ValueError(f"Template checklist header drifted: expected {expected_header}, found {actual_header}.")

    rows: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    contracts = {
        normalize(module): {"module_id": module_id, "module": module, "level": level}
        for module_id, module, level in CHECKLIST_MODULE_CONTRACT
    }
    seen_modules: set[str] = set()
    for row_index, row in enumerate(table.rows[2:], 2):
        cells = row.cells
        if len(cells) < 5:
            raise ValueError(f"Template checklist row {row_index} has fewer than five cells.")
        module, operation, impact, level = [cells[index].text.strip() for index in range(4)]
        if not all((module, operation, impact, level)):
            raise ValueError(f"Template checklist row {row_index} has an empty required cell.")
        module_normalized = normalize(module)
        contract = contracts.get(module_normalized)
        if contract is None:
            raise ValueError(f"Template checklist row {row_index} uses unknown module contract '{module}'.")
        if level != contract["level"]:
            raise ValueError(
                f"Template checklist row {row_index} level '{level}' conflicts with module contract "
                f"'{contract['module_id']}' (expected '{contract['level']}')."
            )
        seen_modules.add(contract["module_id"])
        row_id = stable_row_id(contract["module_id"], operation, impact, level)
        if row_id in seen_ids:
            raise ValueError(f"Template checklist contains duplicate row content for stable id '{row_id}'.")
        seen_ids.add(row_id)
        rows.append(
            {
                "row_id": row_id,
                "row_index": row_index,
                "module_id": contract["module_id"],
                "module": module,
                "operation": operation,
                "impact": impact,
                "level": level,
            }
        )
    expected_modules = {module_id for module_id, _, _ in CHECKLIST_MODULE_CONTRACT}
    missing_modules = sorted(expected_modules - seen_modules)
    if missing_modules:
        raise ValueError(f"Template checklist is missing module contract(s): {', '.join(missing_modules)}.")
    return rows


def resolve_selector(selector: str | dict[str, str], rows: list[dict[str, Any]], field: str) -> dict[str, Any]:
    if isinstance(selector, str):
        criteria = {"operation": selector}
    else:
        criteria = selector
    matches = rows
    for key in ("row_id", "module_id", "module", "operation", "impact", "level"):
        if key in criteria:
            matches = [row for row in matches if normalize(str(row[key])) == normalize(criteria[key])]
    if not matches:
        raise ValueError(f"Checklist selector '{field}' did not match any template row: {criteria}.")
    if len(matches) > 1:
        choices = ", ".join(row["row_id"] for row in matches)
        raise ValueError(
            f"Checklist selector '{field}' is ambiguous; use row_id or module+operation+impact. Matches: {choices}."
        )
    return matches[0]


def checklist_states(data: dict[str, Any], rows: list[dict[str, Any]]) -> dict[str, bool]:
    states: dict[str, bool] = {}

    def register(row: dict[str, Any], involved: bool, field: str) -> None:
        row_id = row["row_id"]
        if row_id in states and states[row_id] is not involved:
            raise ValueError(f"Checklist row '{row_id}' has conflicting decisions (at {field}).")
        states[row_id] = involved

    for index, selector in enumerate(data.get("checklist_items", [])):
        register(resolve_selector(selector, rows, f"checklist_items[{index}]"), True, f"checklist_items[{index}]")
    for index, decision in enumerate(data.get("checklist_decisions", [])):
        selector = {
            key: decision[key]
            for key in ("row_id", "module_id", "module", "operation", "impact", "level")
            if key in decision
        }
        register(
            resolve_selector(selector, rows, f"checklist_decisions[{index}]"),
            decision["involved"],
            f"checklist_decisions[{index}]",
        )
    return states


def update_checklist(doc: Any, data: dict[str, Any]) -> None:
    rows = checklist_rows(doc)
    states = checklist_states(data, rows)
    mark_unmatched = data.get("mark_unmatched_checklist_as_no", False)
    table = unique_table_by_first_cell(doc, CHECKLIST_TITLE)
    for row in rows:
        answer = table.rows[row["row_index"]].cells[4]
        set_cell_text(answer, "")
        if row["row_id"] in states:
            involved = states[row["row_id"]]
            set_cell_text(answer, "是" if involved else "否", force_bold=involved)
        elif mark_unmatched:
            set_cell_text(answer, "否")


def probe_template(doc: Any) -> dict[str, Any]:
    covers = []
    for key, placeholder in COVER_FIELDS.items():
        index, paragraph = unique_paragraph(doc, placeholder)
        covers.append({"key": key, "placeholder": paragraph.text, "paragraph_index": index})

    config_table = unique_table_by_first_cell(doc, "软件信息")
    configs = []
    for key, label in CONFIG_FIELDS.items():
        row, column, target = config_target(config_table, label)
        configs.append(
            {
                "key": key,
                "label": label,
                "target_row": row,
                "target_column": column,
                "current_value": target.text.strip(),
            }
        )

    sections = []
    for key, (heading, next_heading) in SECTION_MAP.items():
        body = body_paragraphs(doc, heading, next_heading)
        sections.append({"key": key, "heading": heading, "body_paragraph_count": len(body)})
    risk_body = body_paragraphs(doc, "3.1风险清单", "3.2风险预案")
    plan_body = body_paragraphs(doc, "变更计划", None)
    rows = checklist_rows(doc)
    return {
        "cover_fields": covers,
        "config_fields": configs,
        "sections": sections
        + [
            {"key": "risks", "heading": "3.1风险清单", "body_paragraph_count": len(risk_body)},
            {"key": "change_plan", "heading": "变更计划", "body_paragraph_count": len(plan_body)},
        ],
        "checklist": {"title": CHECKLIST_TITLE, "row_count": len(rows)},
    }


def update_document(doc: Any, data: dict[str, Any]) -> None:
    probe_template(doc)
    for key, placeholder in COVER_FIELDS.items():
        _, paragraph = unique_paragraph(doc, placeholder)
        if key == "risk_level":
            value = f"风险等级：{data[key]}"
        elif key == "document_date":
            value = f"时间：{data[key]}"
        else:
            value = data[key]
        write_rich_text(paragraph, value, force_bold=key in {"cover_title", "risk_level"})

    config_table = unique_table_by_first_cell(doc, "软件信息")
    for key, label in CONFIG_FIELDS.items():
        _, _, target = config_target(config_table, label)
        set_cell_text(target, data[key])

    update_checklist(doc, data)
    for key, (heading, next_heading) in SECTION_MAP.items():
        replace_section(doc, heading, next_heading, format_items(data[key]))
    replace_section(doc, "3.1风险清单", "3.2风险预案", format_items(data["risks"]))

    if "change_plan" in data:
        plan_lines = format_items(data["change_plan"])
    else:
        plan_lines = [
            f"变更时间：{data['change_time']}",
            f"变更窗口：{data['maintenance_window']}",
            f"变更执行人：{data['executor']}；联系电话：{data['executor_phone']}",
            f"变更监督人：{data['supervisor']}；联系电话：{data['supervisor_phone']}",
        ]
    replace_section(doc, "变更计划", None, plan_lines)


def iter_cell_paragraphs(cell: _Cell) -> Iterable[Paragraph]:
    for paragraph in cell.paragraphs:
        yield paragraph
    for table in cell.tables:
        for row in table.rows:
            for nested_cell in row.cells:
                yield from iter_cell_paragraphs(nested_cell)


def iter_all_paragraphs(doc: Any) -> Iterable[Paragraph]:
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_cell_paragraphs(cell)


def apply_line_spacing(doc: Any) -> None:
    for paragraph in iter_all_paragraphs(doc):
        paragraph.paragraph_format.line_spacing = 1.5


def mark_fields_for_update(doc: Any) -> None:
    update_fields = doc.settings.element.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        doc.settings.element.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    for field in doc.element.iter(qn("w:fldChar")):
        if field.get(qn("w:fldCharType")) == "begin":
            field.set(qn("w:dirty"), "true")


def clean_core_properties(doc: Any, title: str, subject: str) -> None:
    props = doc.core_properties
    props.title = plain_text(title)
    props.subject = subject
    props.author = ""
    props.last_modified_by = ""
    props.revision = 1


def sanitize_xml_part(name: str, data: bytes) -> bytes:
    if not (name.endswith(".xml") or name.endswith(".rels")):
        return data
    root = LET.fromstring(data)
    if name == "docProps/core.xml":
        sensitive_tags = {
            f"{{{DC_NS}}}creator",
            f"{{{CORE_NS}}}lastModifiedBy",
            f"{{{CORE_NS}}}revision",
            f"{{{DCTERMS_NS}}}created",
            f"{{{DCTERMS_NS}}}modified",
        }
        for child in list(root):
            if child.tag in sensitive_tags:
                root.remove(child)
    elif name.endswith(".rels"):
        for relationship in list(root):
            target = relationship.get("Target", "").replace("\\", "/").lower()
            rel_type = relationship.get("Type", "").lower()
            if (
                "customxml" in target
                or target.endswith("docprops/custom.xml")
                or target.endswith("docprops/app.xml")
                or "docprops/thumbnail." in target
                or "custom-properties" in rel_type
                or "extended-properties" in rel_type
                or rel_type.endswith("/metadata/thumbnail")
            ):
                root.remove(relationship)
    elif name == "[Content_Types].xml":
        for child in list(root):
            part_name = child.get("PartName", "").lower()
            if (
                part_name.startswith("/customxml/")
                or part_name.startswith("/docprops/thumbnail.")
                or part_name in {"/docprops/custom.xml", "/docprops/app.xml"}
            ):
                root.remove(child)
    for element in list(root.iter()):
        for attribute in list(element.attrib):
            local_name = LET.QName(attribute).localname.casefold()
            if local_name.startswith("rsid") or local_name in {"durableid", "paraid", "textid"}:
                del element.attrib[attribute]
        if name == "word/settings.xml" and element is not root:
            local_name = LET.QName(element).localname.casefold()
            if local_name in {"docvars", "rsids", "docid"}:
                element.getparent().remove(element)
    if name == "word/settings.xml":
        LET.cleanup_namespaces(root)
        ignorable = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"
        if ignorable in root.attrib:
            declared = {prefix for prefix in root.nsmap if prefix}
            tokens = [token for token in root.get(ignorable, "").split() if token in declared]
            if tokens:
                root.set(ignorable, " ".join(tokens))
            else:
                del root.attrib[ignorable]
    return LET.tostring(root, encoding="utf-8", xml_declaration=True, standalone=True)


def sanitize_package(source: Path, destination: Path) -> None:
    with zipfile.ZipFile(source, "r") as archive_in, zipfile.ZipFile(destination, "w") as archive_out:
        for item in archive_in.infolist():
            normalized = item.filename.replace("\\", "/").lower()
            if (
                normalized in {"docprops/custom.xml", "docprops/app.xml"}
                or normalized.startswith("customxml/")
                or normalized.startswith("docprops/thumbnail.")
            ):
                continue
            archive_out.writestr(item, sanitize_xml_part(item.filename, archive_in.read(item.filename)))


def package_content_types(archive: zipfile.ZipFile) -> tuple[dict[str, str], dict[str, str]]:
    root = ET.fromstring(archive.read("[Content_Types].xml"))
    defaults: dict[str, str] = {}
    overrides: dict[str, str] = {}
    for element in root:
        local_name = element.tag.rsplit("}", 1)[-1]
        if local_name == "Default":
            defaults[element.get("Extension", "").casefold()] = element.get("ContentType", "")
        elif local_name == "Override":
            overrides[element.get("PartName", "").lstrip("/")] = element.get("ContentType", "")
    return defaults, overrides


def part_content_type(part: str, defaults: dict[str, str], overrides: dict[str, str]) -> str:
    if part in overrides:
        return overrides[part]
    extension = part.rsplit(".", 1)[-1].casefold() if "." in part else ""
    return defaults.get(extension, "")


def relationship_part_name(source_part: str | None) -> str:
    if source_part is None:
        return "_rels/.rels"
    directory, filename = posixpath.split(source_part)
    return posixpath.join(directory, "_rels", f"{filename}.rels")


def resolve_relationship_target(source_part: str | None, target: str) -> str:
    decoded = unquote(target).split("#", 1)[0].split("?", 1)[0].replace("\\", "/")
    if decoded.startswith("/"):
        resolved = posixpath.normpath(decoded.lstrip("/"))
    else:
        resolved = posixpath.normpath(posixpath.join(posixpath.dirname(source_part or ""), decoded))
    if resolved == ".." or resolved.startswith("../"):
        raise ValueError(f"DOCX relationship escapes the package root: {target}.")
    return resolved


def read_relationships(archive: zipfile.ZipFile, source_part: str | None) -> list[dict[str, Any]]:
    relationship_part = relationship_part_name(source_part)
    if relationship_part not in archive.namelist():
        return []
    root = ET.fromstring(archive.read(relationship_part))
    relationships = []
    for element in root:
        relationships.append(
            {
                "source": source_part,
                "relationship_part": relationship_part,
                "type": element.get("Type", ""),
                "role": element.get("Type", "").rsplit("/", 1)[-1].casefold(),
                "target": element.get("Target", ""),
                "external": element.get("TargetMode", "").casefold() == "external",
            }
        )
    return relationships


def package_relationship_graph(
    archive: zipfile.ZipFile,
) -> tuple[set[str], dict[str, set[str]], list[dict[str, Any]]]:
    names = set(archive.namelist())
    reachable: set[str] = set()
    roles: dict[str, set[str]] = {}
    external: list[dict[str, Any]] = []
    queue: list[str] = []
    for relationship in read_relationships(archive, None):
        if relationship["external"]:
            external.append(relationship)
            continue
        target = resolve_relationship_target(None, relationship["target"])
        if target not in names:
            raise ValueError(f"DOCX root relationship references missing part '{target}'.")
        roles.setdefault(target, set()).add(relationship["role"])
        queue.append(target)
    while queue:
        part = queue.pop()
        if part in reachable:
            continue
        reachable.add(part)
        for relationship in read_relationships(archive, part):
            if relationship["external"]:
                external.append(relationship)
                continue
            target = resolve_relationship_target(part, relationship["target"])
            if target not in names:
                raise ValueError(f"DOCX relationship from '{part}' references missing part '{target}'.")
            roles.setdefault(target, set()).add(relationship["role"])
            if target not in reachable:
                queue.append(target)
    return reachable, roles, external


def bundled_media_hashes() -> set[str]:
    with zipfile.ZipFile(DEFAULT_TEMPLATE, "r") as archive:
        return {
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.casefold().startswith("word/media/") and not name.endswith("/")
        }


def bundled_structural_attribute_allowlist() -> set[tuple[str, str, str, str]]:
    values: set[tuple[str, str, str, str]] = set()
    with zipfile.ZipFile(DEFAULT_TEMPLATE, "r") as archive:
        for part in archive.namelist():
            if not part.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(part))
            for node in root.iter():
                values.update((part, node.tag, attribute, value) for attribute, value in node.attrib.items())
    return values


def validate_package_no_secrets(archive: zipfile.ZipFile) -> None:
    _, _, external = package_relationship_graph(archive)
    parts = {
        part
        for part in archive.namelist()
        if part.endswith(".xml") or part.endswith(".rels")
    }
    for part in sorted(parts):
        if part not in archive.namelist() or not (part.endswith(".xml") or part.endswith(".rels")):
            continue
        root = ET.fromstring(archive.read(part))
        validate_secret_xml(root, part)
    for relationship in external:
        validate_no_secrets(
            relationship["target"],
            f"{relationship['relationship_part']}:Relationship.Target",
        )


def hidden_style_ids(archive: zipfile.ZipFile, style_parts: Iterable[str]) -> set[str]:
    hidden: set[str] = set()
    based_on: dict[str, str] = {}
    for part in style_parts:
        root = ET.fromstring(archive.read(part))
        for style in (node for node in root.iter() if node.tag.rsplit("}", 1)[-1].casefold() == "style"):
            style_id = style.get(f"{{{W_NS}}}styleId", "")
            if not style_id:
                continue
            if any(
                node.tag.rsplit("}", 1)[-1].casefold() in {"vanish", "webhidden", "specvanish"}
                for node in style.iter()
            ):
                hidden.add(style_id)
            base = next(
                (
                    node.get(f"{{{W_NS}}}val", "")
                    for node in style
                    if node.tag.rsplit("}", 1)[-1].casefold() == "basedon"
                ),
                "",
            )
            if base:
                based_on[style_id] = base
    changed = True
    while changed:
        changed = False
        for style_id, base in based_on.items():
            if base in hidden and style_id not in hidden:
                hidden.add(style_id)
                changed = True
    return hidden


def validate_customer_package(archive: zipfile.ZipFile) -> None:
    defaults, overrides = package_content_types(archive)
    reachable, roles, external = package_relationship_graph(archive)
    allowed_package_parts = reachable | {"[Content_Types].xml", "_rels/.rels"}
    allowed_package_parts.update(
        relationship_part_name(part)
        for part in reachable
        if relationship_part_name(part) in archive.namelist()
    )
    orphan_parts = [
        part
        for part in archive.namelist()
        if not part.endswith("/") and part not in allowed_package_parts
    ]
    if orphan_parts:
        raise ValueError(f"Customer DOCX contains unreferenced or unknown package part(s): {orphan_parts}.")
    all_relationships = [
        relationship
        for source in [None, *sorted(reachable)]
        for relationship in read_relationships(archive, source)
    ]
    for relationship in all_relationships:
        relationship_type = relationship["type"].casefold()
        if relationship["external"]:
            if relationship_type != HYPERLINK_REL_TYPE:
                raise ValueError(
                    f"Customer DOCX contains a forbidden external relationship type in "
                    f"{relationship['relationship_part']}."
                )
        elif relationship_type not in ALLOWED_INTERNAL_REL_TYPES:
            raise ValueError(
                f"Customer DOCX contains an unapproved relationship type in "
                f"{relationship['relationship_part']}."
            )
    forbidden_roles = {
        "afchunk",
        "attachedtemplate",
        "comments",
        "commentsextended",
        "commentsids",
        "control",
        "glossarydocument",
        "oleobject",
        "package",
        "people",
        "vbaproject",
    }
    forbidden_name_prefixes = (
        "customui/",
        "docprops/thumbnail",
        "word/activex/",
        "word/comments",
        "word/embeddings/",
        "word/glossary/",
        "word/people",
        "word/vbaproject",
    )
    for part in archive.namelist():
        lowered = part.casefold()
        content_type = part_content_type(part, defaults, overrides).casefold()
        if lowered.startswith(forbidden_name_prefixes) or any(
            marker in content_type
            for marker in ("activex", "comments", "macroenabled", "oleobject", "vbaproject", "vba-data")
        ):
            raise ValueError(f"Customer DOCX contains forbidden executable, comment, thumbnail, or hidden part '{part}'.")
    for part in reachable:
        if roles.get(part, set()) & forbidden_roles:
            raise ValueError(f"Customer DOCX contains forbidden relationship role(s) for '{part}': {roles[part]}.")
        if not part_content_type(part, defaults, overrides):
            raise ValueError(f"Customer DOCX reachable part '{part}' has no declared content type.")
    for relationship in external:
        role = relationship["role"]
        field = f"{relationship['relationship_part']}:Relationship.Target"
        validate_customer_text(relationship["target"], field)
        if role != "hyperlink":
            raise ValueError(f"Customer DOCX contains forbidden external '{role}' relationship in {field}.")
        target = relationship["target"]
        for _ in range(3):
            decoded_target = unquote(target)
            if decoded_target == target:
                break
            target = decoded_target
        parsed_target = urlsplit(target)
        if parsed_target.scheme.casefold() not in {"http", "https"} or not parsed_target.hostname:
            raise ValueError(f"Customer DOCX external hyperlink uses an unapproved URI scheme in {field}.")

    allowed_media = bundled_media_hashes()
    for part in archive.namelist():
        content_type = part_content_type(part, defaults, overrides).casefold()
        is_media = part.casefold().startswith("word/media/") or content_type.startswith(("image/", "audio/", "video/"))
        if is_media and not part.endswith("/"):
            digest = hashlib.sha256(archive.read(part)).hexdigest()
            if digest not in allowed_media:
                raise ValueError(f"Customer DOCX media part '{part}' is not in the bundled-template hash allowlist.")

    visible_roles = {"officedocument", "header", "footer", "footnotes", "endnotes"}
    visible_parts = {part for part in reachable if roles.get(part, set()) & visible_roles}
    style_parts = {part for part in reachable if "styles" in roles.get(part, set())}
    for part in style_parts:
        root = ET.fromstring(archive.read(part))
        if any(
            node.tag.rsplit("}", 1)[-1].casefold() in {"vanish", "webhidden", "specvanish"}
            for node in root.iter()
        ):
            raise ValueError(f"Customer DOCX styles contain hidden-text style/default markup in {part}.")
    hidden_styles = hidden_style_ids(archive, style_parts)
    if hidden_styles:
        raise ValueError(f"Customer DOCX styles contain hidden-text style(s): {sorted(hidden_styles)}.")
    forbidden_markup = {
        "altchunk",
        "commentrangeend",
        "commentrangestart",
        "commentreference",
        "del",
        "deltext",
        "ins",
        "movefrom",
        "movefromrangeend",
        "movefromrangestart",
        "moveto",
        "movetorangeend",
        "movetorangestart",
        "specvanish",
        "vanish",
        "webhidden",
    }
    for part in sorted(visible_parts):
        root = ET.fromstring(archive.read(part))
        if any(node.tag.rsplit("}", 1)[-1].casefold() in forbidden_markup for node in root.iter()):
            raise ValueError(f"Customer DOCX contains revisions, comments, hidden text, or external chunks in {part}.")
        for node in root.iter():
            if node.tag.rsplit("}", 1)[-1].casefold() in {"pstyle", "rstyle"}:
                if node.get(f"{{{W_NS}}}val", "") in hidden_styles:
                    raise ValueError(f"Customer DOCX uses hidden style '{node.get(f'{{{W_NS}}}val')}' in {part}.")
        validate_customer_xml_text(root, part)
        alt_text = [
            value
            for node in root.iter()
            for attribute, value in node.attrib.items()
            if attribute.rsplit("}", 1)[-1].casefold() in {"descr", "title"}
        ]
        for index, value in enumerate(alt_text):
            validate_customer_text(value, f"{part}:alt[{index}]")
    core = ET.fromstring(archive.read("docProps/core.xml"))
    for index, value in enumerate(node.text or "" for node in core.iter() if node.text):
        validate_customer_text(value, f"docProps/core.xml:text[{index}]")
    structural_allowlist = bundled_structural_attribute_allowlist()
    for part in sorted(allowed_package_parts):
        if not (part.endswith(".xml") or part.endswith(".rels")):
            continue
        root = ET.fromstring(archive.read(part))
        validate_customer_xml_text(root, part)
        for node in root.iter():
            for attribute, value in node.attrib.items():
                if (part, node.tag, attribute, value) not in structural_allowlist:
                    validate_customer_text(value, f"{part}:@{attribute.rsplit('}', 1)[-1]}")
    for source in [None, *sorted(reachable)]:
        for relationship in read_relationships(archive, source):
            validate_customer_text(
                relationship["target"],
                f"{relationship['relationship_part']}:Relationship.Target",
            )


def validate_package(path: Path, audience: str) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        corrupt = archive.testzip()
        if corrupt:
            raise ValueError(f"Generated DOCX contains a corrupt ZIP member: {corrupt}.")
        names = {name.replace("\\", "/").lower() for name in archive.namelist()}
        if (
            names & {"docprops/custom.xml", "docprops/app.xml"}
            or any(name.startswith("customxml/") for name in names)
            or any(name.startswith("docprops/thumbnail.") for name in names)
        ):
            raise ValueError("Generated DOCX still contains identity-bearing property parts.")
        settings = ET.fromstring(archive.read("word/settings.xml"))
        if any(
            element.tag.rsplit("}", 1)[-1].casefold() in {"docvars", "rsids", "docid"}
            for element in settings.iter()
        ):
            raise ValueError("Generated DOCX still contains document variables or document-tracking identifiers.")
        settings_bytes = archive.read("word/settings.xml").lower()
        if any(marker in settings_bytes for marker in (b"commondata", b"userid", b"hdid", b"wpscustom")):
            raise ValueError("Generated DOCX settings still contain WPS identity metadata.")
        for name in (item for item in archive.namelist() if item.endswith(".xml")):
            root = ET.fromstring(archive.read(name))
            for element in root.iter():
                if any(
                    attribute.rsplit("}", 1)[-1].casefold().startswith("rsid")
                    or attribute.rsplit("}", 1)[-1].casefold() in {"durableid", "paraid", "textid"}
                    for attribute in element.attrib
                ):
                    raise ValueError(f"Generated DOCX XML still contains paragraph/session tracking metadata in {name}.")
        validate_package_no_secrets(archive)
        if audience == "customer":
            validate_customer_package(archive)
    Document(str(path))


def resolved_path(path: str | Path) -> Path:
    return Path(path).expanduser().resolve()


def same_path(left: Path, right: Path) -> bool:
    return os.path.normcase(str(left)) == os.path.normcase(str(right))


def validate_paths(input_path: Path, template: Path, output: Path) -> None:
    if same_path(output, input_path):
        raise ValueError("Output path must not be the input JSON path.")
    if same_path(output, template):
        raise ValueError("Output path must not overwrite the template.")
    if output.suffix.lower() != ".docx":
        raise ValueError("Output path must use the .docx extension.")


def atomic_save(doc: Any, output: Path, audience: str) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    rendered: Path | None = None
    sanitized: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent, prefix=f".{output.name}.", suffix=".rendered", delete=False) as handle:
            rendered = Path(handle.name)
        with tempfile.NamedTemporaryFile(dir=output.parent, prefix=f".{output.name}.", suffix=".sanitized", delete=False) as handle:
            sanitized = Path(handle.name)
        doc.save(str(rendered))
        sanitize_package(rendered, sanitized)
        validate_package(sanitized, audience)
        os.replace(sanitized, output)
        sanitized = None
    finally:
        for temporary in (rendered, sanitized):
            if temporary is not None:
                temporary.unlink(missing_ok=True)


def generate(data: dict[str, Any], template: Path, output: Path) -> None:
    validated = validate_input(data)
    if not template.is_file():
        raise FileNotFoundError(f"Template not found: {template}")
    doc = Document(str(template))
    update_document(doc, validated)
    apply_line_spacing(doc)
    mark_fields_for_update(doc)
    clean_core_properties(doc, validated["change_title"], "ZStack change proposal")
    atomic_save(doc, output, validated["audience"])


def reject_duplicate_object(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Input JSON contains duplicate key '{key}'.")
        result[key] = value
    return result


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as handle:
        return validate_input(json.load(handle, object_pairs_hook=reject_duplicate_object))


def print_template(template: Path) -> None:
    if not template.is_file():
        raise FileNotFoundError(f"Template not found: {template}")
    doc = Document(str(template))
    summary = probe_template(doc)
    summary["template"] = str(template)
    json.dump(summary, sys.stdout, ensure_ascii=True, indent=2)
    sys.stdout.write("\n")


def print_checklist(template: Path) -> None:
    if not template.is_file():
        raise FileNotFoundError(f"Template not found: {template}")
    doc = Document(str(template))
    probe_template(doc)
    json.dump(checklist_rows(doc), sys.stdout, ensure_ascii=True, indent=2)
    sys.stdout.write("\n")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_json", nargs="?", help="UTF-8 or UTF-8-BOM JSON input file.")
    parser.add_argument("--out", help="Output DOCX path.")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Template DOCX path.")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--print-checklist", action="store_true", help="Probe and print checklist rows as JSON.")
    mode.add_argument("--print-template", action="store_true", help="Probe and print writable template fields as JSON.")
    args = parser.parse_args(argv)

    try:
        template = resolved_path(args.template)
        if args.print_checklist or args.print_template:
            if args.input_json or args.out:
                raise ValueError("Print modes cannot be combined with input_json or --out.")
            if args.print_checklist:
                print_checklist(template)
            else:
                print_template(template)
            return 0
        if not args.input_json or not args.out:
            parser.error("input_json and --out are required unless a print mode is used.")
        input_path = resolved_path(args.input_json)
        output = resolved_path(args.out)
        validate_paths(input_path, template, output)
        generate(load_json(input_path), template, output)
        print(f"Generated: {output}")
        return 0
    except (OSError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
